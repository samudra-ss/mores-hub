"""MORES ERP - Excel import/export built on openpyxl."""
import io
import re
from datetime import datetime, date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)
NUM_FMT = "#,##0;[Red](#,##0)"
THIN = Border(bottom=Side(style="thin", color="CCCCCC"))


def _sheet(wb, title, report_title, subtitle=""):
    ws = wb.active if wb.active.max_row == 1 and wb.active.max_column == 1 else wb.create_sheet()
    ws.title = title[:31]
    ws["A1"] = report_title
    ws["A1"].font = TITLE_FONT
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(color="666666", italic=True)
    return ws


def _header_row(ws, row, headers, widths=None):
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def _num(ws, row, col, value, bold=False):
    cell = ws.cell(row=row, column=col, value=round(value or 0, 2))
    cell.number_format = NUM_FMT
    if bold:
        cell.font = BOLD
    return cell


def _to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# --------------------------------------------------------------------------
# Exports
# --------------------------------------------------------------------------

def export_trial_balance(tb, scope_label, period_label):
    wb = Workbook()
    ws = _sheet(wb, "Trial Balance", "Trial Balance — %s" % scope_label, period_label)
    _header_row(ws, 4, ["Code", "Account", "Type", "Debit", "Credit", "Balance"],
                [14, 40, 12, 18, 18, 18])
    r = 5
    # grouped by parent account when available: parent rows bold, children
    # indented; balances still reconcile because only leaf rows carry raw amounts
    display = tb.get("grouped") or [dict(x, level=0, is_group=False) for x in tb["rows"]]
    for row in display:
        is_group = row.get("is_group")
        indent = "    " * row.get("level", 0)
        c1 = ws.cell(row=r, column=1, value=row["code"])
        c2 = ws.cell(row=r, column=2, value=indent + row["name"])
        ws.cell(row=r, column=3, value=row["type"].title())
        _num(ws, r, 4, row["debit"], bold=is_group)
        _num(ws, r, 5, row["credit"], bold=is_group)
        _num(ws, r, 6, row["balance"], bold=is_group)
        if is_group:
            c1.font = BOLD
            c2.font = BOLD
        r += 1
    ws.cell(row=r, column=2, value="TOTAL").font = BOLD
    _num(ws, r, 4, tb["total_debit"], bold=True)
    _num(ws, r, 5, tb["total_credit"], bold=True)
    return _to_bytes(wb)


def export_receivables(aging, scope_label):
    wb = Workbook()
    ws = _sheet(wb, "Piutang", "AR Aging (Piutang) — %s" % scope_label,
                "As of %s" % aging["as_of"])
    _header_row(ws, 4, ["No", "Client", "Invoice", "Invoice Date", "Due Date",
                        "Amount", "Outstanding", "Not Due", "1-30 d", "31-60 d",
                        "61-90 d", "> 90 d", "Days Late", "Status"],
                [5, 28, 16, 13, 13, 16, 16, 14, 12, 12, 12, 12, 10, 20])
    bcol = {"not_due": 8, "d1_30": 9, "d31_60": 10, "d61_90": 11, "d90": 12}
    r = 5
    for i, it in enumerate(aging["items"], 1):
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=it["client"])
        ws.cell(row=r, column=3, value=it["invoice_no"])
        _date_cell(ws, r, 4, it["invoice_date"])
        _date_cell(ws, r, 5, it["due_date"])
        _num(ws, r, 6, it["amount"])
        _num(ws, r, 7, it["outstanding"])
        if it["bucket"]:
            _num(ws, r, bcol[it["bucket"]], it["outstanding"])
        ws.cell(row=r, column=13, value=it["days_overdue"] or 0)
        ws.cell(row=r, column=14, value=it["status_label"])
        r += 1
    ws.cell(row=r, column=2, value="TOTAL").font = BOLD
    _num(ws, r, 6, aging["total_amount"], bold=True)
    _num(ws, r, 7, aging["total_outstanding"], bold=True)
    for b, col in bcol.items():
        _num(ws, r, col, aging["buckets"][b], bold=True)
    r += 3
    ws.cell(row=r, column=2, value="AGING SUMMARY").font = BOLD
    r += 1
    for s in aging["summary"]:
        ws.cell(row=r, column=2, value=s["label"])
        _num(ws, r, 6, s["amount"])
        pc = ws.cell(row=r, column=7, value=(s["pct"] or 0) / 100.0)
        pc.number_format = "0.0%"
        r += 1
    ws.cell(row=r, column=2, value="TOTAL OUTSTANDING").font = BOLD
    _num(ws, r, 6, aging["total_outstanding"], bold=True)
    return _to_bytes(wb)


def export_payables(aging, scope_label):
    wb = Workbook()
    ws = _sheet(wb, "Hutang", "AP Aging (Hutang) — %s" % scope_label,
                "As of %s" % aging["as_of"])
    _header_row(ws, 4, ["No", "Vendor", "Bill", "Bill Date", "Due Date",
                        "Amount", "Outstanding", "Not Due", "1-30 d", "31-60 d",
                        "61-90 d", "> 90 d", "Days Late", "Status"],
                [5, 28, 16, 13, 13, 16, 16, 14, 12, 12, 12, 12, 10, 20])
    bcol = {"not_due": 8, "d1_30": 9, "d31_60": 10, "d61_90": 11, "d90": 12}
    r = 5
    for i, it in enumerate(aging["items"], 1):
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=it["vendor"])
        ws.cell(row=r, column=3, value=it["bill_no"])
        _date_cell(ws, r, 4, it["bill_date"])
        _date_cell(ws, r, 5, it["due_date"])
        _num(ws, r, 6, it["amount"])
        _num(ws, r, 7, it["outstanding"])
        if it["bucket"]:
            _num(ws, r, bcol[it["bucket"]], it["outstanding"])
        ws.cell(row=r, column=13, value=it["days_overdue"] or 0)
        ws.cell(row=r, column=14, value=it["status_label"])
        r += 1
    ws.cell(row=r, column=2, value="TOTAL").font = BOLD
    _num(ws, r, 6, aging["total_amount"], bold=True)
    _num(ws, r, 7, aging["total_outstanding"], bold=True)
    for b, col in bcol.items():
        _num(ws, r, col, aging["buckets"][b], bold=True)
    r += 3
    ws.cell(row=r, column=2, value="AGING SUMMARY").font = BOLD
    r += 1
    for s in aging["summary"]:
        ws.cell(row=r, column=2, value=s["label"])
        _num(ws, r, 6, s["amount"])
        pc = ws.cell(row=r, column=7, value=(s["pct"] or 0) / 100.0)
        pc.number_format = "0.0%"
        r += 1
    ws.cell(row=r, column=2, value="TOTAL OUTSTANDING").font = BOLD
    _num(ws, r, 6, aging["total_outstanding"], bold=True)
    return _to_bytes(wb)


# Wallet / card Excel import — downloadable template + format documentation.
# Column names must match exactly (any column ORDER is fine — lookup is by
# header name); extra columns are ignored.
WALLET_TEMPLATE_COLUMNS = [
    # (header, required, format / accepted values, how it is used)
    ("Transaction Type", "Optional",
     "e.g. PAYMENT · INTERNAL_TRANSFER · CARD_ADD_BALANCE · CARD_REFUND_BALANCE",
     "Internal types (transfers between your own wallets/cards) are detected and left unticked."),
    ("Reference ID", "Recommended", "Any unique text, e.g. WLT-2026-0001",
     "Duplicate detection — a Reference ID already booked is flagged. Blank = auto fingerprint."),
    ("Status", "Optional", "SUCCESS · SETTLED · COMPLETED (anything else = row skipped)",
     "Only successful transactions are imported."),
    ("Category", "Optional",
     "FOOD_AND_BEVERAGE · TRANSPORTATION · EXPEDITION_EXPENSES · OFFICE_SUPPLIES · "
     "SOFTWARE · TELECOMMUNICATION · MISCELLANEOUS",
     "Suggests the expense account: OFFICE_SUPPLIES→6600, SOFTWARE/TELECOMMUNICATION→6300, others→6900."),
    ("Transaction Datetime", "REQUIRED", "YYYY-MM-DD HH:MM:SS or DD/MM/YYYY",
     "The journal entry date."),
    ("Amount", "REQUIRED", "Number. NEGATIVE = money out (spending), positive = money in",
     "Spending is booked: DEBIT the expense account (5000–8000) · CREDIT the Petty Cash / Cash & Bank account."),
    ("Description", "Optional", "Free text", "Journal entry description."),
    ("Account Name", "Optional", "Wallet/account label", "Shown as the source account."),
    ("Card Name", "Optional", "Card label", "Shown with the transaction."),
    ("Recipient Holder Name", "Optional", "Counterparty name", "Merged into the description."),
    ("Notes", "Optional", "Free text", "Merged into the description."),
]

WALLET_TEMPLATE_EXAMPLES = [
    ("PAYMENT", "WLT-2026-0001", "SUCCESS", "FOOD_AND_BEVERAGE",
     "2026-06-10 12:30:00", -150000, "Team lunch", "Ops Wallet", "", "Warung Padang Jaya", ""),
    ("PAYMENT", "WLT-2026-0002", "SUCCESS", "TRANSPORTATION",
     "2026-06-11 09:15:00", -85000, "Ride to client meeting", "Ops Wallet", "Corporate Card", "", ""),
    ("PAYMENT", "WLT-2026-0003", "SETTLED", "OFFICE_SUPPLIES",
     "13/06/2026", -230000, "Printer paper & toner", "Ops Wallet", "", "Toko ATK Sentosa", "Q2 restock"),
    ("INTERNAL_TRANSFER", "WLT-2026-0004", "SUCCESS", "",
     "2026-06-14 10:00:00", 5000000, "Top-up from bank", "Ops Wallet", "", "", "left unticked automatically"),
]


def export_wallet_template():
    """The wallet/card import template: a Transactions sheet ready to fill in
    plus a Format Guide sheet documenting every column and the booking rules."""
    wb = Workbook()
    # Transactions sheet: headers on ROW 1 (exactly like the real wallet
    # export) so the filled-in template parses without any preamble handling
    ws = wb.active
    ws.title = "Transactions"
    headers = [c[0] for c in WALLET_TEMPLATE_COLUMNS]
    _header_row(ws, 1, headers, [20, 16, 12, 22, 21, 14, 28, 14, 14, 22, 20])
    r = 2
    for row in WALLET_TEMPLATE_EXAMPLES:
        for col, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=val)
            if headers[col - 1] == "Amount":
                cell.number_format = NUM_FMT
        r += 1

    gd = _sheet(wb, "Format Guide", "Wallet / Card Excel — Format Guide",
                "Fill the Transactions sheet (delete the example rows) and upload it in "
                "Account Parsing → Wallet / Card Excel. Column ORDER does not matter "
                "(lookup is by header name); extra columns are ignored.")
    _header_row(gd, 4, ["Column", "Required", "Format / accepted values", "How it is used"],
                [22, 14, 52, 62])
    r = 5
    for name, req, fmt, use in WALLET_TEMPLATE_COLUMNS:
        gd.cell(row=r, column=1, value=name).font = BOLD
        gd.cell(row=r, column=2, value=req)
        gd.cell(row=r, column=3, value=fmt)
        gd.cell(row=r, column=4, value=use)
        r += 1
    r += 1
    gd.cell(row=r, column=1, value="BOOKING RULES").font = BOLD
    for rule in (
        "Money OUT (negative Amount): DEBIT the expense account you pick (5000–8000 only) "
        "· CREDIT the Petty Cash / Cash & Bank account selected above the table.",
        "Money IN (positive Amount): DEBIT Petty Cash / Cash & Bank · CREDIT the account you pick.",
        "Internal types (INTERNAL_TRANSFER, CARD_ADD_BALANCE, CARD_REFUND_BALANCE) are "
        "detected as moves between your own wallets and left unticked.",
        "Rows whose Status is not SUCCESS / SETTLED / COMPLETED are skipped.",
        "Re-uploading the same file is safe — rows already booked (same Reference ID) are "
        "flagged as duplicates and unticked.",
    ):
        r += 1
        gd.cell(row=r, column=1, value="•")
        gd.cell(row=r, column=2, value=rule)
        gd.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    return _to_bytes(wb)


def export_pnl(pnl, scope_label, period_label):
    wb = Workbook()
    ws = _sheet(wb, "Profit & Loss", "Profit & Loss — %s" % scope_label, period_label)
    _header_row(ws, 4, ["Code", "Account", "Amount"], [10, 42, 20])
    r = 5
    ws.cell(row=r, column=2, value="REVENUE").font = BOLD
    r += 1
    for row in pnl["revenue"]:
        ws.cell(row=r, column=1, value=row["code"])
        ws.cell(row=r, column=2, value=row["name"])
        _num(ws, r, 3, row["balance"])
        r += 1
    ws.cell(row=r, column=2, value="Total Revenue").font = BOLD
    _num(ws, r, 3, pnl["total_revenue"], bold=True)
    r += 2
    ws.cell(row=r, column=2, value="EXPENSES").font = BOLD
    r += 1
    for row in pnl["expense"]:
        ws.cell(row=r, column=1, value=row["code"])
        ws.cell(row=r, column=2, value=row["name"])
        _num(ws, r, 3, row["balance"])
        r += 1
    ws.cell(row=r, column=2, value="Total Expenses").font = BOLD
    _num(ws, r, 3, pnl["total_expense"], bold=True)
    r += 2
    ws.cell(row=r, column=2, value="NET PROFIT").font = Font(bold=True, size=12)
    _num(ws, r, 3, pnl["net_profit"], bold=True)
    return _to_bytes(wb)


def export_balance_sheet(bs, scope_label, period_label):
    wb = Workbook()
    ws = _sheet(wb, "Balance Sheet", "Balance Sheet — %s" % scope_label, period_label)
    _header_row(ws, 4, ["Code", "Account", "Amount"], [10, 42, 20])
    r = 5
    for section, rows, total in [
        ("ASSETS", bs["assets"], bs["total_assets"]),
        ("LIABILITIES", bs["liabilities"], bs["total_liabilities"]),
        ("EQUITY", bs["equity"], bs["total_equity"]),
    ]:
        ws.cell(row=r, column=2, value=section).font = BOLD
        r += 1
        for row in rows:
            ws.cell(row=r, column=1, value=row["code"])
            ws.cell(row=r, column=2, value=row["name"])
            _num(ws, r, 3, row["balance"])
            r += 1
        ws.cell(row=r, column=2, value="Total %s" % section.title()).font = BOLD
        _num(ws, r, 3, total, bold=True)
        r += 2
    return _to_bytes(wb)


def export_budget_vs_actual(bva, scope_label):
    wb = Workbook()
    ws = _sheet(wb, "Budget vs Realization",
                "Budget vs Realization — %s" % scope_label, "Year %s" % bva["year"])
    _header_row(ws, 4, ["Code", "Account", "Type", "Budget", "Realization", "Variance", "Used %"],
                [10, 36, 12, 18, 18, 18, 10])
    r = 5
    for row in bva["rows"]:
        ws.cell(row=r, column=1, value=row["code"])
        ws.cell(row=r, column=2, value=row["name"])
        ws.cell(row=r, column=3, value=row["type"].title())
        _num(ws, r, 4, row["budget"])
        _num(ws, r, 5, row["actual"])
        _num(ws, r, 6, row["variance"])
        if row["used_pct"] is not None:
            ws.cell(row=r, column=7, value=row["used_pct"] / 100).number_format = "0.0%"
        r += 1
    return _to_bytes(wb)


def export_budget_grid(rows, scope_label, year, project_codes=None):
    """rows: [{code, name, project_id, amounts:[12]}]

    Column 2 is the PROJECT CODE, not the account name, so the file this writes
    can be fed straight back into the budget import. It used to write the name
    there, which meant every re-imported row was rejected with "Project code
    'Salaries & Benefits' not found". The name moved to the end, where it is
    still readable and no longer sits in a column the importer reads positionally.
    """
    project_codes = project_codes or {}
    wb = Workbook()
    ws = _sheet(wb, "Budget %s" % year, "Budget — %s" % scope_label, "Year %s" % year)
    _header_row(ws, 4,
                ["Account Code", "Project Code (optional)"] + MONTHS + ["Total", "Account Name"],
                [14, 20] + [13] * 13 + [32])
    r = 5
    for row in rows:
        ws.cell(row=r, column=1, value=row["code"])
        ws.cell(row=r, column=2, value=project_codes.get(row.get("project_id"), ""))
        for m in range(12):
            _num(ws, r, 3 + m, row["amounts"][m])
        _num(ws, r, 15, sum(row["amounts"]), bold=True)
        ws.cell(row=r, column=16, value=row["name"])
        r += 1
    return _to_bytes(wb)


def export_journals(entries, scope_label, period_label):
    """entries: [{entry_no, date, status, description, reference, lines:[{account_code, account_name, project_code, description, debit, credit}]}]"""
    wb = Workbook()
    ws = _sheet(wb, "Journal Entries", "Journal Entries — %s" % scope_label, period_label)
    _header_row(ws, 4, ["Entry No", "Date", "Status", "Entry Description", "Account Code",
                        "Account Name", "Project", "Line Description", "Debit", "Credit"],
                [16, 12, 9, 32, 12, 28, 12, 30, 17, 17])
    r = 5
    for e in entries:
        for ln in e["lines"]:
            ws.cell(row=r, column=1, value=e["entry_no"])
            _date_cell(ws, r, 2, e["date"])
            ws.cell(row=r, column=3, value=e["status"])
            ws.cell(row=r, column=4, value=e["description"])
            ws.cell(row=r, column=5, value=ln["account_code"])
            ws.cell(row=r, column=6, value=ln["account_name"])
            ws.cell(row=r, column=7, value=ln.get("project_code") or "")
            ws.cell(row=r, column=8, value=ln["description"])
            _num(ws, r, 9, ln["debit"])
            _num(ws, r, 10, ln["credit"])
            r += 1
    return _to_bytes(wb)


def export_coa(accounts, scope_label):
    wb = Workbook()
    ws = _sheet(wb, "Chart of Accounts", "Chart of Accounts — %s" % scope_label)
    _header_row(ws, 4, ["Code", "Name", "Type", "Parent Code", "Intercompany", "Active"],
                [10, 40, 12, 12, 13, 8])
    r = 5
    for a in accounts:
        ws.cell(row=r, column=1, value=a["code"])
        ws.cell(row=r, column=2, value=a["name"])
        ws.cell(row=r, column=3, value=a["type"].title())
        ws.cell(row=r, column=4, value=a["parent_code"] or "")
        ws.cell(row=r, column=5, value="Y" if a["is_intercompany"] else "")
        ws.cell(row=r, column=6, value="Y" if a["is_active"] else "N")
        r += 1
    return _to_bytes(wb)


def export_project_performance(rows, scope_label, year):
    wb = Workbook()
    ws = _sheet(wb, "Project Performance",
                "Project Performance — %s" % scope_label, "Year %s" % year)
    _header_row(ws, 4, ["Company", "Code", "Project", "Status", "Revenue", "Expense",
                        "Profit", "Margin %", "Budget Rev", "Budget Exp"],
                [10, 12, 32, 10, 18, 18, 18, 10, 18, 18])
    r = 5
    for p in rows:
        ws.cell(row=r, column=1, value=p["company"])
        ws.cell(row=r, column=2, value=p["code"])
        ws.cell(row=r, column=3, value=p["name"])
        ws.cell(row=r, column=4, value=p["status"])
        _num(ws, r, 5, p["revenue"])
        _num(ws, r, 6, p["expense"])
        _num(ws, r, 7, p["profit"])
        ws.cell(row=r, column=8, value=p["margin_pct"] / 100).number_format = "0.0%"
        _num(ws, r, 9, p["budget_revenue"])
        _num(ws, r, 10, p["budget_expense"])
        r += 1
    return _to_bytes(wb)


def export_cash_flow(cf, scope_label):
    wb = Workbook()
    ws = _sheet(wb, "Cash Flow", "Cash Flow Analysis — %s" % scope_label, "Year %s" % cf["year"])
    _header_row(ws, 4, ["Month", "Cash In", "Cash Out", "Net", "Ending Balance"],
                [12, 20, 20, 20, 20])
    ws.cell(row=5, column=1, value="Opening").font = BOLD
    _num(ws, 5, 5, cf["opening_balance"], bold=True)
    r = 6
    for m in cf["monthly"]:
        ws.cell(row=r, column=1, value=MONTHS[m["month"] - 1])
        _num(ws, r, 2, m["cash_in"])
        _num(ws, r, 3, m["cash_out"])
        _num(ws, r, 4, m["net"])
        _num(ws, r, 5, m["ending"])
        r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = BOLD
    _num(ws, r, 2, cf["total_in"], bold=True)
    _num(ws, r, 3, cf["total_out"], bold=True)
    _num(ws, r, 4, cf["net_change"], bold=True)
    _num(ws, r, 5, cf["closing_balance"], bold=True)
    r += 2
    ws.cell(row=r, column=1, value="SOURCES OF CASH").font = BOLD
    ws.cell(row=r, column=4, value="USES OF CASH").font = BOLD
    r += 1
    for i in range(max(len(cf["sources"]), len(cf["uses"]))):
        if i < len(cf["sources"]):
            s = cf["sources"][i]
            ws.cell(row=r + i, column=1, value="%s %s" % (s["code"], s["name"]))
            _num(ws, r + i, 2, s["amount"])
        if i < len(cf["uses"]):
            u = cf["uses"][i]
            ws.cell(row=r + i, column=4, value="%s %s" % (u["code"], u["name"]))
            _num(ws, r + i, 5, u["amount"])
    return _to_bytes(wb)


# --------------------------------------------------------------------------
# Templates
# --------------------------------------------------------------------------

def template_journals():
    wb = Workbook()
    ws = _sheet(wb, "Journals", "Journal Import Template",
                "One row per line. Rows sharing the same Entry Ref form one entry; debits must equal credits.")
    _header_row(ws, 4, ["Entry Ref", "Date (YYYY-MM-DD)", "Description", "Account Code",
                        "Line Description", "Debit", "Credit", "Project Code"],
                [12, 18, 32, 14, 30, 16, 16, 14])
    sample = [
        ("INV-001", "2026-06-01", "Customer invoice", "1200", "Invoice #123", 50000000, 0, "PRJ-APP"),
        ("INV-001", "2026-06-01", "Customer invoice", "4100", "Invoice #123", 0, 50000000, "PRJ-APP"),
        ("PAY-001", "2026-06-05", "Office supplies", "6600", "Stationery", 1500000, 0, ""),
        ("PAY-001", "2026-06-05", "Office supplies", "1120", "Bank payment", 0, 1500000, ""),
    ]
    for i, row in enumerate(sample):
        for j, v in enumerate(row, start=1):
            ws.cell(row=5 + i, column=j, value=v)
    return _to_bytes(wb)


def template_coa():
    wb = Workbook()
    ws = _sheet(wb, "COA", "Chart of Accounts Import Template",
                "Type must be one of: Asset, Liability, Equity, Revenue, Expense")
    _header_row(ws, 4, ["Code", "Name", "Type", "Parent Code", "Intercompany (Y/N)"],
                [10, 40, 12, 12, 16])
    ws.append([])
    for row in [("1130", "Petty Cash Branch", "Asset", "1100", ""),
                ("6800", "Travel & Entertainment", "Expense", "6000", "")]:
        ws.append(row)
    return _to_bytes(wb)


def template_budget(year, projects=None, project_code=None):
    """The budget import template.

    Column B is the project code: blank is a company-level line, a code puts the
    row on that project. The codes that exist are listed on a second sheet,
    because "Project Code (optional)" with nothing to copy from is exactly how a
    project NAME ends up typed into a column that reads codes.
    """
    wb = Workbook()
    ws = _sheet(wb, "Budget", "Budget Import Template",
                "Amounts per month for year %s. Account Code must exist in the company COA. "
                "Leave Project Code empty for a company-level line." % year)
    _header_row(ws, 4, ["Account Code", "Project Code (optional)"] + MONTHS, [14, 22] + [13] * 12)
    ws.cell(row=5, column=1, value="6100")
    if project_code:
        ws.cell(row=5, column=2, value=project_code)
    for m in range(12):
        ws.cell(row=5, column=3 + m, value=500000000)
    if projects:
        wp = wb.create_sheet("Project list")
        wp["A1"] = "Project codes you can put in column B"
        wp["A1"].font = TITLE_FONT
        _header_row(wp, 3, ["Project Code", "Project Name"], [22, 52])
        for i, p in enumerate(projects):
            wp.cell(row=4 + i, column=1, value=p.get("code"))
            wp.cell(row=4 + i, column=2, value=p.get("name"))
    return _to_bytes(wb)


# --------------------------------------------------------------------------
# Imports
# --------------------------------------------------------------------------

def upsert_budget_month(conn, company_id, account_id, project_id, year, month, amount):
    """Write a MONTH total into a week-grain budget without destroying the shape.

    Excel is month grain and the budget is week grain, so an import has to decide
    where in the month the money goes. Dumping it all in week 1 would silently
    undo whatever week placement the user set in Budget Center - export, re-import,
    and a deliberately W4-heavy month is a W1 spike again. So: if the month
    already has a shape, keep the shape and rescale it to the new total; only a
    month with nothing in it falls back to week 1.
    """
    existing = conn.execute(
        "SELECT week, amount FROM budgets WHERE company_id=? AND account_id=?"
        " AND project_id IS ? AND year=? AND month=?",
        (company_id, account_id, project_id, year, month)).fetchall()
    shape = [(r[0], r[1]) for r in existing if r[1]]
    total_now = sum(a for _, a in shape)
    amount = round(amount or 0, 2)
    if not shape or not total_now:
        for wk, _ in existing:
            if wk != 1:
                conn.execute(
                    "DELETE FROM budgets WHERE company_id=? AND account_id=?"
                    " AND project_id IS ? AND year=? AND month=? AND week=?",
                    (company_id, account_id, project_id, year, month, wk))
        upsert_budget(conn, company_id, account_id, project_id, year, month, amount, week=1)
        return
    # rescale, and give the last bucket the rounding remainder so the month is exact
    allocated = 0.0
    for i, (wk, a) in enumerate(shape):
        part = (round(amount - allocated, 2) if i == len(shape) - 1
                else round(amount * a / total_now, 2))
        allocated = round(allocated + part, 2)
        upsert_budget(conn, company_id, account_id, project_id, year, month, part, week=wk)


def upsert_budget(conn, company_id, account_id, project_id, year, month, amount, week=1):
    """NULL-safe budget upsert into ONE week bucket."""
    cur = conn.execute(
        "UPDATE budgets SET amount=? WHERE company_id=? AND account_id=?"
        " AND project_id IS ? AND year=? AND month=? AND week=?",
        (round(amount, 2), company_id, account_id, project_id, year, month, week),
    )
    if cur.rowcount == 0:
        conn.execute(
            "INSERT INTO budgets (company_id, account_id, project_id, year, month, week, amount)"
            " VALUES (?,?,?,?,?,?,?)",
            (company_id, account_id, project_id, year, month, week, round(amount, 2)),
        )


def _date_cell(ws, row, col, value):
    """Write an ISO date as a REAL Excel date displayed DD/MM/YYYY, so it sorts
    and filters as a date and reads the way the rest of the app does."""
    cell = ws.cell(row=row, column=col)
    text = str(value or "")[:10]
    try:
        cell.value = datetime.strptime(text, "%Y-%m-%d")
        cell.number_format = "DD/MM/YYYY"
    except ValueError:
        cell.value = value
    return cell


def _iso_date(text):
    """DD/MM/YYYY (how people type it) or YYYY-MM-DD (how a real Excel date
    arrives) -> YYYY-MM-DD, or None when it is neither."""
    text = (text or "").strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _cell_str(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _cell_num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return 0.0


def _find_header(ws, first_col_name):
    for row in ws.iter_rows(min_row=1, max_row=10):
        if _cell_str(row[0].value).lower().startswith(first_col_name.lower()):
            return row[0].row
    return None


def import_journals(conn, company_id, stream, created_by):
    """Returns (created_entries, errors)."""
    wb = load_workbook(stream, data_only=True)
    ws = wb.active
    hdr = _find_header(ws, "Entry Ref")
    if not hdr:
        return 0, ["Header row not found — first column must be 'Entry Ref'. Use the template."]
    accounts = {r["code"]: r["id"] for r in conn.execute(
        "SELECT id, code FROM accounts WHERE company_id=? AND is_active=1", (company_id,))}
    projects = {r["code"]: r["id"] for r in conn.execute(
        "SELECT id, code FROM projects WHERE company_id=?", (company_id,))}

    groups, errors = {}, []
    for row in ws.iter_rows(min_row=hdr + 1, values_only=True):
        if not row or not _cell_str(row[0]):
            continue
        ref = _cell_str(row[0])
        entry_date = _iso_date(_cell_str(row[1]))
        if not entry_date:
            errors.append("Entry %s: invalid date '%s' (use DD/MM/YYYY)" % (ref, _cell_str(row[1])))
            continue
        code = _cell_str(row[3])
        if code not in accounts:
            errors.append("Entry %s: account code '%s' not found in this company" % (ref, code))
            continue
        pcode = _cell_str(row[7]) if len(row) > 7 else ""
        if pcode and pcode not in projects:
            errors.append("Entry %s: project code '%s' not found" % (ref, pcode))
            continue
        g = groups.setdefault(ref, {"date": entry_date, "description": _cell_str(row[2]), "lines": []})
        g["lines"].append({
            "account_id": accounts[code],
            "description": _cell_str(row[4]),
            "debit": _cell_num(row[5]),
            "credit": _cell_num(row[6]),
            "project_id": projects.get(pcode),
        })

    if errors:
        return 0, errors

    created = 0
    for ref, g in groups.items():
        total_d = round(sum(l["debit"] for l in g["lines"]), 2)
        total_c = round(sum(l["credit"] for l in g["lines"]), 2)
        if abs(total_d - total_c) > 0.01:
            errors.append("Entry %s is not balanced (debit %.2f vs credit %.2f) — skipped" % (ref, total_d, total_c))
            continue
        if total_d == 0:
            errors.append("Entry %s has zero amount — skipped" % ref)
            continue
        n = conn.execute(
            "SELECT COUNT(*)+1 FROM journal_entries WHERE company_id=?", (company_id,)
        ).fetchone()[0]
        entry_no = "IMP-%s-%05d" % (g["date"][:7].replace("-", ""), n)
        cur = conn.execute(
            "INSERT INTO journal_entries (company_id, entry_no, date, description, reference, status, source, created_by)"
            " VALUES (?,?,?,?,?,'draft','excel',?)",
            (company_id, entry_no, g["date"], g["description"], ref, created_by),
        )
        for l in g["lines"]:
            conn.execute(
                "INSERT INTO journal_lines (entry_id, account_id, project_id, description, debit, credit)"
                " VALUES (?,?,?,?,?,?)",
                (cur.lastrowid, l["account_id"], l["project_id"], l["description"],
                 round(l["debit"], 2), round(l["credit"], 2)),
            )
        created += 1
    conn.commit()
    return created, errors


def import_coa(conn, company_id, stream):
    wb = load_workbook(stream, data_only=True)
    ws = wb.active
    hdr = _find_header(ws, "Code")
    if not hdr:
        return 0, 0, ["Header row not found — first column must be 'Code'. Use the template."]
    valid_types = {"asset", "liability", "equity", "revenue", "expense"}
    created = updated = 0
    errors = []
    for row in ws.iter_rows(min_row=hdr + 1, values_only=True):
        if not row or not _cell_str(row[0]):
            continue
        code, name = _cell_str(row[0]), _cell_str(row[1])
        typ = _cell_str(row[2]).lower()
        parent = _cell_str(row[3]) if len(row) > 3 else ""
        if not parent and "-" in code:
            parent = code.rsplit("-", 1)[0]  # derivative: 5100-01-01 -> 5100-01
        ic = 1 if (len(row) > 4 and _cell_str(row[4]).upper() in ("Y", "YES", "1")) else 0
        if typ not in valid_types:
            errors.append("Account %s: invalid type '%s'" % (code, typ))
            continue
        if code.count("-") > 2:
            errors.append("Account %s: maximum 3 levels (e.g. 5100-01-01)" % code)
            continue
        existing = conn.execute(
            "SELECT id FROM accounts WHERE company_id=? AND code=?", (company_id, code)
        ).fetchone()
        if existing:
            conn.execute(
                "UPDATE accounts SET name=?, type=?, parent_code=?, is_intercompany=? WHERE id=?",
                (name, typ, parent or None, ic, existing["id"]))
            updated += 1
        else:
            conn.execute(
                "INSERT INTO accounts (company_id, code, name, type, parent_code, is_intercompany)"
                " VALUES (?,?,?,?,?,?)",
                (company_id, code, name, typ, parent or None, ic))
            created += 1
    conn.commit()
    return created, updated, errors


def import_budget(conn, company_id, year, stream, project_id=None):
    """Read a budget workbook -> (rows saved, errors, rows per project).

    `project_id` means "this file is for that project": a blank Project Code
    lands there instead of at company level, and a row naming a DIFFERENT
    project is refused. Both halves matter - a per-project import that quietly
    writes company-level rows looks like it worked and shows up nowhere.
    """
    wb = load_workbook(stream, data_only=True)
    ws = wb.active
    hdr = _find_header(ws, "Account Code")
    if not hdr:
        return 0, ["Header row not found — first column must be 'Account Code'. Use the template."], {}
    accounts = {r["code"]: r["id"] for r in conn.execute(
        "SELECT id, code FROM accounts WHERE company_id=?", (company_id,))}
    projects = {r["code"]: r["id"] for r in conn.execute(
        "SELECT id, code FROM projects WHERE company_id=?", (company_id,))}
    codes = {v: k for k, v in projects.items()}
    saved, errors, by_project = 0, [], {}
    for n, row in enumerate(ws.iter_rows(min_row=hdr + 1, values_only=True), start=hdr + 1):
        if not row or not _cell_str(row[0]):
            continue
        code = _cell_str(row[0])
        if code not in accounts:
            errors.append("Row %d: account code '%s' is not in this company — row skipped" % (n, code))
            continue
        pcode = _cell_str(row[1]) if len(row) > 1 else ""
        if pcode and pcode not in projects:
            errors.append("Row %d: project code '%s' is not in this company — row skipped" % (n, pcode))
            continue
        pid = projects[pcode] if pcode else project_id
        if project_id and pid != project_id:
            errors.append("Row %d: this import is for project %s, but the row says %s — row skipped"
                          % (n, codes.get(project_id, project_id), pcode))
            continue
        for m in range(12):
            amount = _cell_num(row[2 + m]) if len(row) > 2 + m else 0.0
            upsert_budget_month(conn, company_id, accounts[code], pid, year, m + 1, amount)
        saved += 1
        where = codes.get(pid) if pid else "company level"
        by_project[where] = by_project.get(where, 0) + 1
    conn.commit()
    return saved, errors, by_project


# --------------------------------------------------------------------------
# SBU model workbook - set up (or back up) a whole SBU from one Excel file
# --------------------------------------------------------------------------

# key, label, kind, note. The key sits in a hidden column so a label can be
# reworded (or translated) inside the file without breaking the import.
SBU_FIELDS = [
    ("name", "SBU name", "text", "Required."),
    ("code", "Code", "text", "e.g. SBU-01 - blank keeps the current code"),
    ("company", "Company code", "text", "The owning company, e.g. SBR"),
    ("stage", "Stage", "stage", "idea / build / launch / growth / sunset"),
    ("launch_month", "Launch month", "month", "MM/YYYY"),
    ("target_year", "Must be profitable by (year)", "int", "e.g. 2027"),
    ("actual_through", "Actual through", "month", "MM/YYYY - blank = the last month the ledger has"),
    ("horizon_end", "Model until", "month", "MM/YYYY - blank = December of the target year"),
    ("burn_budget", "Burn budget (Rp)", "money", "What the company is prepared to lose before it pays back"),
    ("notes", "Notes", "text", ""),
    ("arpu", "Price per active user / month (Rp)", "money", "ARPU"),
    ("arpu_growth_annual", "Price rise per year", "pct", "A percentage"),
    ("revenue_start_month", "Subscription revenue starts", "month", "MM/YYYY - blank = the launch month"),
    ("start_users", "Users already on board then", "number", ""),
    ("new_users_first", "New users in the first month", "number", ""),
    ("new_users_growth", "Growth in new users / month", "pct", "A percentage"),
    ("churn_monthly", "Churn / month", "pct", "A percentage of the opening base lost each month"),
    ("server_base", "Server base cost / month (Rp)", "money", "Runs whether anyone logs in or not"),
    ("server_per_user", "Server cost per active user / month (Rp)", "money", ""),
    ("server_step_users", "Add a capacity tier every N users", "int", "0 = no tiers"),
    ("server_step_cost", "Each tier costs / month (Rp)", "money", ""),
]
_SBU_ASSUMPTIONS = {"arpu", "arpu_growth_annual", "revenue_start_month", "start_users", "new_users_first",
                    "new_users_growth", "churn_monthly", "server_base", "server_per_user",
                    "server_step_users", "server_step_cost"}
_PCT_FMT = "0.00%"


def _month_out(v):
    v = str(v or "")
    return "%s/%s" % (v[5:7], v[:4]) if re.match(r"^\d{4}-\d{2}", v) else None


def _sbu_month_in(v):
    """MM/YYYY as typed, a real Excel date, YYYY-MM or DD/MM/YYYY -> YYYY-MM."""
    if v in (None, ""):
        return None
    if isinstance(v, (datetime, date)):
        return "%04d-%02d" % (v.year, v.month)
    s = str(v).strip()
    for pattern, yi, mi in ((r"^(\d{1,2})[/-](\d{4})$", 2, 1),
                            (r"^(\d{4})-(\d{1,2})(?:-\d{1,2})?$", 1, 2),
                            (r"^\d{1,2}/(\d{1,2})/(\d{4})$", 2, 1)):
        m = re.match(pattern, s)
        if m and 1 <= int(m.group(mi)) <= 12:
            return "%s-%02d" % (m.group(yi), int(m.group(mi)))
    raise ValueError("'%s' is not a month - write it as MM/YYYY" % s)


def _sbu_money_in(v):
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).replace("Rp", "").replace("rp", "").replace(" ", "").strip()
    if re.match(r"^-?\d{1,3}(\.\d{3})+(,\d+)?$", s):        # 1.200.000,50
        s = s.replace(".", "").replace(",", ".")
    elif re.match(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$", s):      # 1,200,000.50
        s = s.replace(",", "")
    elif "," in s and "." not in s:                          # 2,5
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        raise ValueError("'%s' is not an amount" % v)


def _sbu_pct_in(cell):
    """The workbook holds real Excel percentages (0.02 shown as 2%). A plain
    cell holding 2, or the text '2%', means 2% as well."""
    v = getattr(cell, "value", cell)
    if v in (None, ""):
        return 0.0
    if isinstance(v, str):
        s = v.strip()
        return _sbu_money_in(s[:-1] if s.endswith("%") else s) / 100.0
    fmt = getattr(cell, "number_format", "") or ""
    return float(v) if "%" in fmt else float(v) / 100.0


def _sbu_int_in(v):
    if v in (None, ""):
        return 0
    try:
        return int(round(_sbu_money_in(v)))
    except ValueError:
        raise ValueError("'%s' is not a whole number" % v)


def _sbu_category_in(v, categories):
    s = _cell_str(v).lower()
    if not s:
        return "other"
    for key, label in categories:
        if s in (key, label.lower()):
            return key
    raise ValueError("unknown category '%s' - use %s" % (_cell_str(v), " / ".join(k for k, _ in categories)))


def _note(ws, ref, text):
    ws[ref] = text
    ws[ref].font = Font(color="666666", italic=True)


def export_sbu_workbook(sbu, company_code, linked_projects, all_projects, categories, stages):
    """One workbook that IS the SBU model: settings and drivers, cost lines,
    one-offs, CAPEX and linked projects. Download it, edit it, import it back.
    An empty `sbu` gives the blank template."""
    wb = Workbook()
    a = sbu.get("assumptions") or {}
    ws = _sheet(wb, "SBU", "SBU model - %s" % (sbu.get("name") or "new SBU"),
                "Fill in the Value column. Months are MM/YYYY, percentages are percentages. "
                "Importing the file replaces the SBU's whole model with what it says.")
    _header_row(ws, 4, ["Field", "Value", "Notes", "key"], [42, 26, 66, 4])
    stage_row = None
    for i, (key, label, kind, note) in enumerate(SBU_FIELDS):
        r = 5 + i
        ws.cell(row=r, column=1, value=label).font = BOLD
        val = company_code if key == "company" else (a.get(key) if key in _SBU_ASSUMPTIONS else sbu.get(key))
        cell = ws.cell(row=r, column=2)
        if kind == "month":
            cell.value = _month_out(val)
        elif kind == "pct":
            cell.value = float(val or 0)
            cell.number_format = _PCT_FMT
        elif kind == "money":
            cell.value = float(val or 0)
            cell.number_format = NUM_FMT
        elif kind in ("int", "number"):
            cell.value = val if val not in (None, "") else 0
        else:
            cell.value = val or None
        if key == "stage":
            stage_row = r
        ws.cell(row=r, column=3, value=note or None).font = Font(color="666666", italic=True)
        ws.cell(row=r, column=4, value=key).font = Font(color="BBBBBB")
    ws.column_dimensions["D"].hidden = True
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(stages), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B%d" % stage_row)

    cats = [k for k, _ in categories]
    lines = sbu.get("lines") or []
    wl = wb.create_sheet("Cost lines")
    wl["A1"] = "Forecast cost lines"
    wl["A1"].font = TITLE_FONT
    _note(wl, "A2", "Category: %s. Basis: per month or per working day. A blank From means the line "
                    "runs from the first forecast month." % " / ".join(cats))
    _header_row(wl, 4, ["Line", "Category", "Amount (Rp)", "Basis", "From (MM/YYYY)", "Until (MM/YYYY)",
                        "Rise per year"], [38, 14, 18, 18, 16, 16, 14])
    for i, ln in enumerate(lines):
        r = 5 + i
        wl.cell(row=r, column=1, value=ln.get("label") or None)
        wl.cell(row=r, column=2, value=ln.get("category") or "other")
        _num(wl, r, 3, ln.get("monthly_amount"))
        wl.cell(row=r, column=4, value="per working day" if ln.get("basis") == "per_working_day" else "per month")
        wl.cell(row=r, column=5, value=_month_out(ln.get("start_month")))
        wl.cell(row=r, column=6, value=_month_out(ln.get("end_month")))
        wl.cell(row=r, column=7, value=float(ln.get("escalation_annual") or 0)).number_format = _PCT_FMT
    for r in range(5 + len(lines), 201):
        wl.cell(row=r, column=7).number_format = _PCT_FMT   # a typed 5 becomes 5%, as Excel does
    for col, options in (("B", cats), ("D", ["per month", "per working day"])):
        v = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True)
        wl.add_data_validation(v)
        v.add("%s5:%s500" % (col, col))

    wo = wb.create_sheet("One-offs")
    wo["A1"] = "One-off items"
    wo["A1"].font = TITLE_FONT
    _note(wo, "A2", "Direction: in (money coming in, e.g. a termin) or out. Months that are already "
                    "actual are ignored - the ledger decides what happened in those.")
    _header_row(wo, 4, ["Month (MM/YYYY)", "Direction", "What", "Amount (Rp)"], [18, 12, 42, 18])
    for i, o in enumerate(sbu.get("oneoffs") or []):
        r = 5 + i
        wo.cell(row=r, column=1, value=_month_out(o.get("month")))
        wo.cell(row=r, column=2, value="in" if o.get("flow") == "in" else "out")
        wo.cell(row=r, column=3, value=o.get("label") or None)
        _num(wo, r, 4, o.get("amount"))
    v = DataValidation(type="list", formula1='"in,out"', allow_blank=True)
    wo.add_data_validation(v)
    v.add("B5:B500")

    wc = wb.create_sheet("CAPEX")
    wc["A1"] = "CAPEX register"
    wc["A1"].font = TITLE_FONT
    _note(wc, "A2", "Cash when bought, then depreciated straight-line over its life.")
    _header_row(wc, 4, ["Month bought (MM/YYYY)", "Asset", "Amount (Rp)", "Life (months)"], [22, 42, 18, 14])
    for i, cx in enumerate(sbu.get("capex") or []):
        r = 5 + i
        wc.cell(row=r, column=1, value=_month_out(cx.get("month")))
        wc.cell(row=r, column=2, value=cx.get("label") or None)
        _num(wc, r, 3, cx.get("amount"))
        wc.cell(row=r, column=4, value=int(cx.get("life_months") or 48))

    wp = wb.create_sheet("Projects")
    wp["A1"] = "Ledger projects this SBU spends from"
    wp["A1"].font = TITLE_FONT
    _note(wp, "A2", "Actual months come from posted journal lines on these projects, in whichever "
                    "company's books they were booked. Copy codes from the 'Project list' sheet.")
    _header_row(wp, 4, ["Project code", "Company code (needed only if the code exists in two companies)",
                        "Project name (reference only)"], [18, 30, 46])
    for i, p in enumerate(linked_projects or []):
        r = 5 + i
        wp.cell(row=r, column=1, value=p.get("code") or None)
        wp.cell(row=r, column=2, value=p.get("company_code") or None)
        wp.cell(row=r, column=3, value=p.get("name") or None)

    wr = wb.create_sheet("Project list")
    wr["A1"] = "Every project you can link"
    wr["A1"].font = TITLE_FONT
    _header_row(wr, 3, ["Project code", "Company code", "Project name"], [18, 16, 52])
    for i, p in enumerate(all_projects or []):
        r = 4 + i
        wr.cell(row=r, column=1, value=p.get("code") or None)
        wr.cell(row=r, column=2, value=p.get("company_code") or None)
        wr.cell(row=r, column=3, value=p.get("name") or None)
    wb.active = 0
    return _to_bytes(wb)


def parse_sbu_workbook(stream, categories, stages):
    """Read an SBU model workbook -> (payload, errors).

    The payload has the shape PUT /api/products/<id> takes, plus company_code and
    project_refs for the server to resolve. Nothing is written here, and the
    server applies nothing when any error is reported: half an imported model is
    worse than none, because it looks finished.
    """
    try:
        wb = load_workbook(stream, data_only=True)
    except Exception:
        return None, ["That file could not be opened as an Excel workbook (.xlsx)."]
    sheets = {w.title.strip().lower(): w for w in wb.worksheets}
    ws = sheets.get("sbu")
    if ws is None:
        return None, ["This is not an SBU model workbook - it has no sheet named 'SBU'. "
                      "Download the template from the SBU page."]
    errors, payload = [], {"assumptions": {}}
    kinds = {k: kind for k, _, kind, _ in SBU_FIELDS}
    by_label = {label.lower(): k for k, label, _, _ in SBU_FIELDS}
    for row in ws.iter_rows(min_row=5):
        if not row:
            continue
        label = _cell_str(row[0].value)
        key = _cell_str(row[3].value) if len(row) > 3 else ""
        if key not in kinds:
            key = by_label.get(label.lower())
        if not key:
            continue
        cell = row[1] if len(row) > 1 else None
        raw = cell.value if cell is not None else None
        kind = kinds[key]
        try:
            if kind == "month":
                val = _sbu_month_in(raw)
            elif kind == "pct":
                val = _sbu_pct_in(cell)
            elif kind in ("money", "number"):
                val = _sbu_money_in(raw)
            elif kind == "int":
                val = _sbu_int_in(raw)
            elif kind == "stage":
                val = _cell_str(raw).lower() or "build"
                if val not in stages:
                    raise ValueError("stage must be one of %s" % " / ".join(stages))
            else:
                val = _cell_str(raw)
        except ValueError as e:
            errors.append("SBU sheet, %s: %s" % (label or key, e))
            continue
        if key == "company":
            payload["company_code"] = val
        elif key in _SBU_ASSUMPTIONS:
            payload["assumptions"][key] = val
        else:
            payload[key] = val
    if not payload.get("name"):
        errors.append("SBU sheet: 'SBU name' is required.")
    for key in ("code", "target_year"):            # blank keeps what the SBU already has
        if not payload.get(key):
            payload.pop(key, None)

    def rows_of(name):
        s = sheets.get(name)
        return None if s is None else list(enumerate(s.iter_rows(min_row=5), start=5))

    def at(r, i):
        return r[i].value if i < len(r) else None

    rs = rows_of("cost lines")
    if rs is not None:
        lines = []
        for n, r in rs:
            label, amount = _cell_str(at(r, 0)), at(r, 2)
            if not label and amount in (None, ""):
                continue
            try:
                lines.append({
                    "label": label or "Cost line",
                    "category": _sbu_category_in(at(r, 1), categories),
                    "monthly_amount": _sbu_money_in(amount),
                    "basis": "per_working_day" if "working" in _cell_str(at(r, 3)).lower() else "fixed",
                    "start_month": _sbu_month_in(at(r, 4)),
                    "end_month": _sbu_month_in(at(r, 5)),
                    "escalation_annual": _sbu_pct_in(r[6] if len(r) > 6 else None),
                })
            except ValueError as e:
                errors.append("Cost lines row %d: %s" % (n, e))
        payload["lines"] = lines

    rs = rows_of("one-offs")
    if rs is not None:
        out = []
        for n, r in rs:
            m, amount = at(r, 0), at(r, 3)
            if m in (None, "") and amount in (None, ""):
                continue
            try:
                month = _sbu_month_in(m)
                if not month:
                    raise ValueError("a month is required")
                d = _cell_str(at(r, 1)).lower()
                if d in ("in", "money in", "masuk"):
                    flow = "in"
                elif d in ("", "out", "money out", "keluar"):
                    flow = "out"
                else:
                    raise ValueError("direction must be in or out")
                out.append({"month": month, "flow": flow, "label": _cell_str(at(r, 2)),
                            "amount": _sbu_money_in(amount)})
            except ValueError as e:
                errors.append("One-offs row %d: %s" % (n, e))
        payload["oneoffs"] = out

    rs = rows_of("capex")
    if rs is not None:
        out = []
        for n, r in rs:
            m, amount = at(r, 0), at(r, 2)
            if m in (None, "") and amount in (None, ""):
                continue
            try:
                month = _sbu_month_in(m)
                if not month:
                    raise ValueError("a month is required")
                out.append({"month": month, "label": _cell_str(at(r, 1)), "amount": _sbu_money_in(amount),
                            "life_months": _sbu_int_in(at(r, 3)) or 48})
            except ValueError as e:
                errors.append("CAPEX row %d: %s" % (n, e))
        payload["capex"] = out

    rs = rows_of("projects")
    if rs is not None:
        payload["project_refs"] = [(n, _cell_str(at(r, 0)), _cell_str(at(r, 1)))
                                   for n, r in rs if _cell_str(at(r, 0))]
    return payload, errors


# --------------------------------------------------------------------------
# Weekly cashflow projection workbook -> the Oracle
#
# The shape is the one finance keeps by hand (and the one in the Q4 workbook):
# a row of month headers, a W1..W4 row under it, and below each week a column of
# labels with the amount in the column next to it - positive money in, negative
# money out. Nothing here writes: it reads the sheet and says what it found.
# --------------------------------------------------------------------------

CF_OB_LABELS = {"ob", "opening", "opening balance", "open balance", "saldo awal", "kas awal"}
CF_MONTH_WORDS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "mei": 5, "jun": 6,
                  "jul": 7, "aug": 8, "agu": 8, "ags": 8, "sep": 9, "oct": 10, "okt": 10,
                  "nov": 11, "dec": 12, "des": 12}
CF_MON3 = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


def _cf_week_no(v):
    """'W1', 'w 1', 'Minggu 2', 'Week 3' -> the week number, else None."""
    s = _cell_str(v).lower().replace(".", " ").replace("-", " ")
    m = re.match(r"^(?:w|week|minggu|mg)\s*([1-5])$", s)
    return int(m.group(1)) if m else None


def _cf_month(v, default_year):
    """A month header cell -> (year, month), else None."""
    if isinstance(v, (datetime, date)):
        return v.year, v.month
    s = _cell_str(v).lower().strip()
    if not s:
        return None
    m = re.match(r"^(\d{1,2})[/-](\d{4})$", s)                  # 09/2026
    if m and 1 <= int(m.group(1)) <= 12:
        return int(m.group(2)), int(m.group(1))
    m = re.match(r"^(\d{4})-(\d{1,2})", s)                      # 2026-09
    if m and 1 <= int(m.group(2)) <= 12:
        return int(m.group(1)), int(m.group(2))
    m = re.match(r"^([a-z]{3,})[^a-z0-9]*(\d{2,4})?$", s)       # Sep / September 26 / Des-2026
    if m and m.group(1)[:3] in CF_MONTH_WORDS:
        y = m.group(2)
        year = default_year if not y else (2000 + int(y) if len(y) == 2 else int(y))
        return year, CF_MONTH_WORDS[m.group(1)[:3]]
    return None


def _cf_money(v):
    return "{:,.0f}".format(v or 0)


def _cf_num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def parse_cashflow_workbook(stream, default_year=None):
    """Read a weekly cashflow projection sheet -> (plan, errors).

    A blank row ends the plan: everything under it is somebody's scratch working,
    not the forecast, and importing that would invent money. What was skipped is
    reported rather than dropped quietly.
    """
    default_year = int(default_year or datetime.now().year)
    try:
        wb = load_workbook(stream, data_only=True)
    except Exception:
        return None, ["That file could not be opened as an Excel workbook (.xlsx)."]

    # 1) the week row: whichever row carries the most W1..W4 headers
    best = None
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
            found = [(c.column, _cf_week_no(c.value)) for c in row]
            found = [(col, w) for col, w in found if w]
            if len(found) >= 2 and (best is None or len(found) > len(best[2])):
                best = (ws, row[0].row, sorted(found))
    if best is None:
        return None, ["No week columns found. The sheet needs a row of W1 W2 W3 W4 headers, "
                      "with the month written above them."]
    ws, wrow, weeks = best

    # 2) the month headers above it; the nearest header at or left of a week column wins
    months = []
    for r in range(max(1, wrow - 4), wrow):
        found = [(c.column, _cf_month(c.value, default_year)) for c in ws[r]]
        found = [(col, ym) for col, ym in found if ym]
        if len(found) > len(months):
            months = sorted(found)
    if not months:
        return None, ["No month headers found above the W1..W4 row. Write the month over each "
                      "block - a date like 01/09/2026, or a name like September."]

    def month_for(col):
        hit = months[0][1]
        for mcol, ym in months:
            if mcol <= col:
                hit = ym
        return hit

    buckets, order = {}, []

    def bucket(col, week):
        y, m = month_for(col)
        key = (y, m, week)
        if key not in buckets:
            buckets[key] = {"year": y, "month": m, "week": week,
                            "label": "W%d %s %d" % (week, CF_MON3[m - 1], y),
                            "items": [], "opening": None, "closing": None}
            order.append(key)
        return buckets[key]

    # 3) the rows under the week headers, until the first fully blank one
    warnings, ignored, stopped_at = [], [], None
    for r in range(wrow + 1, ws.max_row + 1):
        cells = {col: (ws.cell(row=r, column=col).value, ws.cell(row=r, column=col + 1).value)
                 for col, _ in weeks}
        if all(not _cell_str(lab) and _cf_num(amt) is None for lab, amt in cells.values()):
            if any(b["items"] for b in buckets.values()):
                stopped_at = r
                break
            continue
        is_ob_row = any(_cell_str(lab).lower() in CF_OB_LABELS for lab, _ in cells.values())
        for col, week in weeks:
            lab, amt = cells[col]
            text, num = _cell_str(lab), _cf_num(amt)
            b = bucket(col, week)
            if is_ob_row:
                if num is not None:
                    b["opening"] = num
                continue
            if text.lower() in CF_OB_LABELS:
                continue
            if not text:
                if num is not None:
                    b["closing"] = num          # a bare number is a running balance
                continue
            if num is None or not num:
                warnings.append("%s (row %d): '%s' has no amount beside it, so it was skipped."
                                % (b["label"], r, text))
                continue
            b["items"].append({"label": text, "amount": abs(num),
                               "flow": "in" if num > 0 else "out", "row": r})

    if stopped_at:
        for r in range(stopped_at, ws.max_row + 1):
            for col, _ in weeks:
                text = _cell_str(ws.cell(row=r, column=col).value)
                num = _cf_num(ws.cell(row=r, column=col + 1).value)
                if text and num:
                    ignored.append("row %d '%s'" % (r, text))

    out_weeks = [buckets[k] for k in order if buckets[k]["items"] or buckets[k]["opening"] is not None]
    items = [i for b in out_weeks for i in b["items"]]
    if not items:
        return None, ["No cashflow rows found under the week headers. Each row needs a label with "
                      "its amount in the column next to it (negative = money out)."]

    years = sorted({b["year"] for b in out_weeks if b["items"]})
    if len(years) > 1:
        warnings.append("The file covers %s. The Oracle reads one year at a time, so the rows for "
                        "the other year are stored but only appear when you consult that year."
                        % " and ".join(str(y) for y in years))
    if ignored:
        warnings.append("%d row(s) below the first blank row were left out - a blank row ends the "
                        "plan: %s%s" % (len(ignored), ", ".join(ignored[:4]),
                                        " …" if len(ignored) > 4 else ""))
    # The sheet carries its own opening balance per week, so it can be asked
    # whether it agrees with itself. A file that does not add up week to week is
    # usually a mis-read column, and that must be visible before it is imported.
    checks = []
    with_open = [b for b in out_weeks if b["opening"] is not None]
    for a, b in zip(with_open, with_open[1:]):
        moved = sum(i["amount"] if i["flow"] == "in" else -i["amount"] for i in a["items"])
        expect = round(a["opening"] + moved, 2)
        if abs(expect - b["opening"]) > 1:
            checks.append(
                "%s opens at %s and its own rows come to %s, but the next week opens at %s "
                "(a gap of %s)." % (a["label"], _cf_money(a["opening"]), _cf_money(expect),
                                    _cf_money(b["opening"]), _cf_money(b["opening"] - expect)))

    plan = {
        "sheet": ws.title, "year": years[0],
        "checks": checks, "reconciled": len(with_open) > 1 and not checks,
        "weeks": out_weeks,
        "months": sorted({(b["year"], b["month"]) for b in out_weeks if b["items"]}),
        "count": len(items),
        "total_in": round(sum(i["amount"] for i in items if i["flow"] == "in"), 2),
        "total_out": round(sum(i["amount"] for i in items if i["flow"] == "out"), 2),
        "opening": next((b["opening"] for b in out_weeks if b["opening"] is not None), None),
        "warnings": warnings, "ignored": ignored,
    }
    return plan, []


# --------------------------------------------------------------------------
# The Oracle's own plan template - built FROM the database
#
# The hand-kept cash sheet above says "25.000.000" and nothing else. This one
# says WHICH account and WHICH project, because every code is a drop-down fed by
# a reference sheet and the budget already in the system is written out as rows.
# A line like "5100-01 Direct Labor / Consultant Fees, project NX-01, 25.000.000
# out in W2 of December 2026" is then picked from lists, not typed from memory.
# --------------------------------------------------------------------------

PLAN_SHEET = "Cash plan"
PLAN_HEADERS = ["Account Code", "Project Code", "Year", "Month", "Week", "Direction",
                "Amount (Rp)", "Certainty", "Cash class", "Note", "Account Name (reference)"]


def has_plan_sheet(stream):
    """Is this the system's own plan template, or somebody's own cash sheet?"""
    try:
        wb = load_workbook(stream, data_only=True, read_only=True)
    except Exception:
        return False
    try:
        return any(w.strip().lower() == PLAN_SHEET.lower() for w in wb.sheetnames)
    finally:
        wb.close()


def _dv(ws, ref, formula):
    v = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(v)
    v.add(ref)


def export_plan_template(year, scope_label, accounts, projects, rows, project_stats=None,
                         last_row=400):
    """rows: [{account_code, account_name, project_code, year, month, week, flow,
               amount, certainty, cf_class, note}] - the budget as it stands."""
    wb = Workbook()
    ws = _sheet(wb, PLAN_SHEET, "Oracle cash plan - %s" % scope_label,
                "One row per movement. Pick the account and project from the drop-downs; the rows "
                "below are the budget already in the system for %d. Example: 5100-01 Direct Labor / "
                "Consultant Fees, project NX-01, 25.000.000, out, W2 of December %d." % (year, year))
    _header_row(ws, 4, PLAN_HEADERS, [14, 16, 8, 10, 8, 11, 18, 14, 13, 34, 32])
    r = 5
    for row in rows or []:
        ws.cell(row=r, column=1, value=row.get("account_code"))
        ws.cell(row=r, column=2, value=row.get("project_code") or None)
        ws.cell(row=r, column=3, value=row.get("year") or year)
        ws.cell(row=r, column=4, value=MONTHS[(row.get("month") or 1) - 1])
        ws.cell(row=r, column=5, value="W%d" % (row.get("week") or 1))
        ws.cell(row=r, column=6, value=row.get("flow") or "out")
        _num(ws, r, 7, row.get("amount"))
        ws.cell(row=r, column=8, value=row.get("certainty") or "planned")
        ws.cell(row=r, column=9, value=row.get("cf_class") or None)
        ws.cell(row=r, column=10, value=row.get("note") or None)
        ws.cell(row=r, column=11, value=row.get("account_name") or None).font = Font(color="888888")
        r += 1
    for rr in range(r, last_row + 1):                 # empty rows keep the drop-downs and format
        ws.cell(row=rr, column=7).number_format = NUM_FMT   # format only - a written 0 is a row
    ws.freeze_panes = "A5"                            # the header stays while the budget scrolls
    ws.auto_filter.ref = "A4:K%d" % max(r - 1, 5)

    wa = wb.create_sheet("Accounts")
    wa["A1"] = "Accounts in this company"
    wa["A1"].font = TITLE_FONT
    _header_row(wa, 3, ["Account Code", "Account Name", "Type", "Cash class"], [16, 40, 12, 13])
    for i, a in enumerate(accounts or []):
        wa.cell(row=4 + i, column=1, value=a.get("code"))
        wa.cell(row=4 + i, column=2, value=a.get("name"))
        wa.cell(row=4 + i, column=3, value=a.get("type"))
        wa.cell(row=4 + i, column=4, value=a.get("cash_flow_class") or "")
    wp = wb.create_sheet("Projects")
    wp["A1"] = "Projects in this company"
    wp["A1"].font = TITLE_FONT
    _header_row(wp, 3, ["Project Code", "Project Name"], [18, 48])
    for i, p in enumerate(projects or []):
        wp.cell(row=4 + i, column=1, value=p.get("code"))
        wp.cell(row=4 + i, column=2, value=p.get("name"))

    rng = "$A$4:$A$%d" % (3 + max(len(accounts or []), 1))
    _dv(ws, "A5:A%d" % last_row, "=Accounts!%s" % rng)
    _dv(ws, "B5:B%d" % last_row, "=Projects!$A$4:$A$%d" % (3 + max(len(projects or []), 1)))
    _dv(ws, "D5:D%d" % last_row, '"%s"' % ",".join(MONTHS))
    _dv(ws, "E5:E%d" % last_row, '"W1,W2,W3,W4"')
    _dv(ws, "F5:F%d" % last_row, '"in,out"')
    _dv(ws, "H5:H%d" % last_row, '"committed,planned,expected,speculative"')
    _dv(ws, "I5:I%d" % last_row, '"operating,investing,financing,noncash"')

    if project_stats is not None:
        wr = wb.create_sheet("Project revenue & budget")
        wr["A1"] = "What each project is budgeted to earn and spend, and what the ledger has booked"
        wr["A1"].font = TITLE_FONT
        _note(wr, "A2", "Budget is this year's budget; realization is posted journal lines tagged to "
                        "the project, in whichever company booked them.")
        _header_row(wr, 4, ["Project Code", "Project Name", "Budget revenue", "Realized revenue",
                            "Budget cost", "Realized cost", "Budget gross profit"],
                    [16, 36, 18, 18, 18, 18, 20])
        for i, p in enumerate(project_stats):
            rr = 5 + i
            wr.cell(row=rr, column=1, value=p.get("code"))
            wr.cell(row=rr, column=2, value=p.get("name"))
            _num(wr, rr, 3, p.get("budget_revenue"))
            _num(wr, rr, 4, p.get("actual_revenue"))
            _num(wr, rr, 5, p.get("budget_cost"))
            _num(wr, rr, 6, p.get("actual_cost"))
            _num(wr, rr, 7, (p.get("budget_revenue") or 0) - (p.get("budget_cost") or 0), bold=True)
    wb.active = 0
    return _to_bytes(wb)


def _plan_month(v):
    """'Dec', 'December', 12, '2026-12' -> 12."""
    if isinstance(v, (datetime, date)):
        return v.month
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        m = int(v)
        if 1 <= m <= 12:
            return m
        raise ValueError("month %s is not between 1 and 12" % v)
    s = _cell_str(v).lower()
    if not s:
        raise ValueError("a month is required")
    if s.isdigit() and 1 <= int(s) <= 12:
        return int(s)
    m = re.match(r"^(\d{4})-(\d{1,2})", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return int(m.group(2))
    if s[:3] in CF_MONTH_WORDS:
        return CF_MONTH_WORDS[s[:3]]
    raise ValueError("'%s' is not a month" % _cell_str(v))


def parse_plan_template(stream, year_hint=None):
    """Read the plan template -> (plan, errors). Nothing is written.

    Codes are not resolved here - the caller knows which company it is importing
    into, and resolving them anywhere else risks booking a row against another
    company's account with the same code.
    """
    year_hint = int(year_hint or datetime.now().year)
    try:
        wb = load_workbook(stream, data_only=True)
    except Exception:
        return None, ["That file could not be opened as an Excel workbook (.xlsx)."]
    ws = next((w for w in wb.worksheets if w.title.strip().lower() == PLAN_SHEET.lower()), None)
    if ws is None:
        return None, ["This is not a plan template - it has no '%s' sheet. Download the plan "
                      "template from the Oracle." % PLAN_SHEET]
    hdr = _find_header(ws, "Account Code")
    if not hdr:
        return None, ["The '%s' sheet needs a header row starting with 'Account Code'." % PLAN_SHEET]

    rows, errors, blank = [], [], 0
    for n, row in enumerate(ws.iter_rows(min_row=hdr + 1, values_only=True), start=hdr + 1):
        def at(i):
            return row[i] if row and i < len(row) else None
        code = _cell_str(at(0))
        try:
            amount = _sbu_money_in(at(6))
        except ValueError as e:
            errors.append("Row %d: %s" % (n, e))
            continue
        if not amount:              # no money on the row: nothing to plan, whatever else it says
            blank += 1
            continue
        if not code:
            errors.append("Row %d: an amount with no account code." % n)
            continue
        try:
            year = _sbu_int_in(at(2)) or year_hint
            month = _plan_month(at(3))
            week = _cf_week_no(at(4))
            if week is None:
                try:
                    week = _sbu_int_in(at(4))
                except ValueError:
                    week = None
            if not 1 <= (week or 0) <= 4:
                raise ValueError("'%s' is not a week - use W1, W2, W3 or W4" % _cell_str(at(4)))
            flow = _cell_str(at(5)).lower()
            if flow not in ("", "in", "out"):
                raise ValueError("direction must be in or out")
            if not flow:
                flow = "out" if amount < 0 else ""        # a minus sign still means money out
            rows.append({
                "row": n, "account_code": code, "project_code": _cell_str(at(1)),
                "year": year, "month": month, "week": week, "flow": flow,
                "amount": abs(amount), "certainty": _cell_str(at(7)).lower(),
                "cf_class": _cell_str(at(8)).lower(), "note": _cell_str(at(9)),
            })
        except ValueError as e:
            errors.append("Row %d: %s" % (n, e))
    if not rows and not errors:
        return None, ["The '%s' sheet has no rows with an amount." % PLAN_SHEET]
    return {"sheet": ws.title, "rows": rows, "count": len(rows),
            "years": sorted({r["year"] for r in rows}), "warnings": []}, errors


# ---------------------------------------------------------------------------
# SBU finance report (v1.10) - the six sheets of "Template Finance Report
# NX-Sentimind", filled from sbu_report.build(). Variances, YTD and totals are
# live formulas, as in the template, so finance can keep working in the file.
# ---------------------------------------------------------------------------

_RP = "#,##0;[Red]-#,##0"
_NAVY, _BLUE, _TEAL, _GREEN, _AMBER = "0D1B2A", "1A5276", "0E6655", "1A7A4A", "D68910"
_DEEP, _RED, _CYAN, _SEC, _GREY = "1A252F", "B03A2E", "028090", "1B4F72", "566573"
_STATUS_ID = {"safe": ("🟢 AMAN", "E9F7EF", _GREEN), "watch": ("🟡 PERHATIAN", "FEF9E7", _AMBER),
              "critical": ("🔴 KRITIS", "FADBD8", _RED)}
_UNIT_FMT = {"rp": "#,##0", "pct": "0.00%", "months": '0.0 "Bulan"', "x": '0.00"x"', "days": '0 "Hari"'}


def _fill(hex6):
    return PatternFill("solid", fgColor="FF" + hex6)


def _box(ws, ref, value=None, fill=None, color="FFFFFF", bold=False, size=9, fmt=None,
         align="left", wrap=False):
    first = ref.split(":")[0]
    if ":" in ref:
        ws.merge_cells(ref)
    c = ws[first]
    if value is not None:
        c.value = value
    if fill:
        for row in ws[ref] if ":" in ref else [[c]]:
            for x in row:
                x.fill = _fill(fill)
    c.font = Font(color="FF" + color, bold=bold, size=size)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fmt:
        c.number_format = fmt
    return c


def _limit_text(v, unit, direction):
    if v is None:
        return "-"
    sign = "≥ " if direction == "high" else "≤ "
    if unit == "pct":
        return sign + ("%.2f%%" % (v * 100)).replace(".", ",")
    if unit == "rp":
        return sign + "{:,.0f}".format(v).replace(",", ".")
    if unit == "months":
        return sign + ("%g Bulan" % v)
    if unit == "days":
        return sign + ("%g Hari" % v)
    return sign + ("%gx" % v)


def export_sbu_report(rep):
    wb = Workbook()
    title = (rep["title"] or "SBU").upper()
    # ================================================================ DASHBOARD
    ws = wb.active
    ws.title = "DASHBOARD"
    for col, w in (("A", 2.4), ("B", 15.6), ("C", 9), ("D", 9), ("E", 2.4), ("F", 15.6), ("G", 9),
                   ("H", 9), ("I", 2.4), ("J", 15.6), ("K", 9), ("L", 9), ("M", 2.4), ("N", 18.5),
                   ("O", 9), ("P", 9), ("Q", 2.4)):
        ws.column_dimensions[col].width = w
    _box(ws, "B1:M2", "%s  |  FINANCIAL DASHBOARD EKSEKUTIF" % title, _NAVY, bold=True, size=14)
    _box(ws, "N1:P2", "Update: %s  |  %s" % (datetime.strptime(rep["generated"], "%Y-%m-%d").strftime("%d/%m/%Y"),
                                             rep["quarter"]), _NAVY, size=9, align="right")
    spots = [("B", "D"), ("F", "H"), ("J", "L"), ("N", "P")]
    colors = [(_NAVY, "02C39A"), (_BLUE, "02C39A"), (_TEAL, "FFFFFF"), (_GREEN, "FFFFFF"),
              (_AMBER, "FFFFFF"), (_DEEP, "02C39A"), (_RED, "FFFFFF"), (_CYAN, "FFFFFF")]
    for i, tile in enumerate(rep["tiles"]):
        a, b = spots[i % 4]
        top = 5 if i < 4 else 12
        fill, vcol = colors[i]
        _box(ws, "%s%d:%s%d" % (a, top, b, top), tile["label"], fill, color="BDC3C7", size=8)
        v = tile["value"]
        if tile["unit"] == "months":
            v = ("~ %s Bulan" % ("%.0f" % v if v is not None else "∞")) if v is not None or tile.get("status") == "safe" else "-"
        elif v is None:
            v = "-"
        _box(ws, "%s%d:%s%d" % (a, top + 1, b, top + 4), v, fill, color=vcol, bold=True, size=18,
             fmt=_UNIT_FMT.get(tile["unit"]) if isinstance(v, (int, float)) else None, align="center")
        _box(ws, "%s%d:%s%d" % (a, top + 5, b, top + 5), tile["sub"], fill, color="BDC3C7", size=8)
    _box(ws, "B19:P19", "🚦  INDIKATOR KESEHATAN KEUANGAN — %s %s" % (title, rep["quarter"]), _NAVY,
         bold=True, size=10)
    for ref, text in (("B20:C20", "Indikator"), ("D20:F20", "Nilai Aktual"), ("G20:H20", "Target 🟢"),
                      ("I20:J20", "Batas Min 🟡"), ("K20:L20", "Status"),
                      ("M20:P20", "Catatan & Rekomendasi Direksi")):
        _box(ws, ref, text, _NAVY, bold=True)
    for n, ind in enumerate(rep["dashboard_indicators"]):
        r = 21 + n
        st, bg, fg = _STATUS_ID[ind["status"]]
        _box(ws, "B%d:C%d" % (r, r), ind["label"], "F0F4F8", color=_DEEP, bold=True)
        val = ind["value"] if ind["value"] is not None else "-"
        _box(ws, "D%d:F%d" % (r, r), val, bg, color=fg, bold=True,
             fmt=_UNIT_FMT.get(ind["unit"]) if ind["value"] is not None else None)
        _box(ws, "G%d:H%d" % (r, r), _limit_text(ind["safe"], ind["unit"], ind["direction"]), "F0F4F8", color=_GREY)
        _box(ws, "I%d:J%d" % (r, r), _limit_text(ind["attention"], ind["unit"], ind["direction"]), "F0F4F8", color="7F8C8D")
        _box(ws, "K%d:L%d" % (r, r), st, bg, color=fg, bold=True)
        _box(ws, "M%d:P%d" % (r, r), ind["note"], "F0F4F8", color=_GREY, size=8, wrap=True)
        ws.row_dimensions[r].height = 30
    _box(ws, "B35:E35", "🟢 AMAN = Dalam target", "E9F7EF", color=_GREEN, bold=True, size=8)
    _box(ws, "F35:I35", "🟡 PERHATIAN = Perlu monitoring", "FEF9E7", color=_AMBER, bold=True, size=8)
    _box(ws, "J35:M35", "🔴 KRITIS = Tindakan segera", "FADBD8", color=_RED, bold=True, size=8)
    _box(ws, "N35:P35", "Perbarui setiap tutup bulan", "F0F4F8", color=_GREY, size=8)

    # ================================================================== P&L YTD
    ws = wb.create_sheet("P&L YTD")
    T = rep["through"]
    months = rep["months"]
    last = get_column_letter(3 + 36 + 3)          # AP
    _box(ws, "A1:%s1" % last, "LAPORAN LABA RUGI — BUDGET VS AKTUAL %d  |  %s" % (rep["year"], rep["title"]),
         _NAVY, bold=True, size=13)
    _box(ws, "A2:%s2" % last, "Semua nilai dalam IDR (Rupiah). Kolom Aktual dari jurnal terposting pada proyek SBU, "
         "ditutup sampai %s." % (rep["through_label"] or "-"), None, color=_GREY, size=9)
    _box(ws, "A3:B4", "KATEGORI", _NAVY, bold=True)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 36
    mcol = lambda m, k: 3 + (m - 1) * 3 + k            # k: 0 budget, 1 actual, 2 variance
    for m in range(1, 13):
        c0 = mcol(m, 0)
        _box(ws, "%s3:%s3" % (get_column_letter(c0), get_column_letter(c0 + 2)), months[m - 1], _SEC,
             bold=True, align="center")
        for k, lbl in enumerate(("Budget", "Aktual", "Variance")):
            _box(ws, "%s4" % get_column_letter(c0 + k), lbl, _GREY, bold=True, size=8, align="center")
            ws.column_dimensions[get_column_letter(c0 + k)].width = 16
    ycol = 3 + 36
    _box(ws, "%s3:%s3" % (get_column_letter(ycol), get_column_letter(ycol + 2)),
         "YTD %s–%s" % (months[0], months[T - 1] if T else "-"), _SEC, bold=True, align="center")
    for k, lbl in enumerate(("Bud YTD", "Akt YTD", "Var YTD")):
        _box(ws, "%s4" % get_column_letter(ycol + k), lbl, _GREY, bold=True, size=8, align="center")
        ws.column_dimensions[get_column_letter(ycol + k)].width = 17
    _box(ws, "%s3" % get_column_letter(ycol + 3), "FULL YEAR", _SEC, bold=True, align="center")
    _box(ws, "%s4" % get_column_letter(ycol + 3), "Budget", _GREY, bold=True, size=8, align="center")
    ws.column_dimensions[get_column_letter(ycol + 3)].width = 17
    ws.freeze_panes = "C5"

    lines = {ln["key"]: ln for ln in rep["pnl"]["lines"]}
    row = [5]

    def section(text, fill):
        r = row[0]
        _box(ws, "A%d:%s%d" % (r, last, r), text, fill, bold=True)
        row[0] += 1

    def sub(text):
        _box(ws, "B%d" % row[0], text, None, color=_DEEP, bold=True)
        row[0] += 1

    def line(key, is_revenue):
        r, ln = row[0], lines[key]
        _box(ws, "B%d" % r, "    " + ln["label"], None, color=_DEEP)
        for m in range(1, 13):
            b, a, v = (get_column_letter(mcol(m, k)) for k in range(3))
            _num(ws, r, mcol(m, 0), ln["budget"][m - 1])
            if m <= T:
                _num(ws, r, mcol(m, 1), ln["actual"][m - 1])
                ws["%s%d" % (v, r)] = ("=%s%d-%s%d" % (a, r, b, r)) if is_revenue else ("=%s%d-%s%d" % (b, r, a, r))
                ws["%s%d" % (v, r)].number_format = NUM_FMT
        _ytd_cells(r, is_revenue)
        row[0] += 1
        return r

    def _ytd_cells(r, is_revenue, bold=False):
        bc = ["%s%d" % (get_column_letter(mcol(m, 0)), r) for m in range(1, T + 1)]
        ac = ["%s%d" % (get_column_letter(mcol(m, 1)), r) for m in range(1, T + 1)]
        fc = ["%s%d" % (get_column_letter(mcol(m, 0)), r) for m in range(1, 13)]
        Y = [get_column_letter(ycol + k) for k in range(4)]
        ws["%s%d" % (Y[0], r)] = ("=" + "+".join(bc)) if bc else 0
        ws["%s%d" % (Y[1], r)] = ("=" + "+".join(ac)) if ac else 0
        ws["%s%d" % (Y[2], r)] = ("=%s%d-%s%d" % (Y[1], r, Y[0], r)) if is_revenue else ("=%s%d-%s%d" % (Y[0], r, Y[1], r))
        ws["%s%d" % (Y[3], r)] = "=" + "+".join(fc)
        for y in Y:
            ws["%s%d" % (y, r)].number_format = NUM_FMT
            if bold:
                ws["%s%d" % (y, r)].font = BOLD

    def total(label, formula_of, fill=_NAVY, is_revenue=True):
        """formula_of(column letter) -> the right-hand side for a budget or actual
        column. A variance is never summed or subtracted from other variances: it
        is this row's own actual against its own budget, so a profit row reads
        actual minus budget and a cost row budget minus actual."""
        r = row[0]
        _box(ws, "B%d" % r, label, fill, bold=True)
        cols = [mcol(m, k) for m in range(1, 13) for k in range(3)] + [ycol + k for k in range(4)]
        for c in cols:
            L = get_column_letter(c)
            m = (c - 3) // 3 + 1 if c < ycol else None
            if m and m > T and (c - 3) % 3 in (1, 2):
                continue
            is_var = (c < ycol and (c - 3) % 3 == 2) or c == ycol + 2
            if is_var:
                b, a = get_column_letter(c - 2), get_column_letter(c - 1)
                ws["%s%d" % (L, r)] = ("=%s%d-%s%d" % (a, r, b, r)) if is_revenue else ("=%s%d-%s%d" % (b, r, a, r))
            else:
                ws["%s%d" % (L, r)] = "=" + formula_of(L)
            ws["%s%d" % (L, r)].number_format = NUM_FMT
            ws["%s%d" % (L, r)].font = BOLD
            ws["%s%d" % (L, r)].fill = _fill("EAF2F8")
        row[0] += 1
        return r

    section("PENDAPATAN (REVENUE)", _CYAN)
    rv = [line("rev_saas", True), line("rev_project", True)]
    r_rev = total("TOTAL PENDAPATAN", lambda L: "SUM(%s%d:%s%d)" % (L, rv[0], L, rv[-1]))
    row[0] += 1
    section("HARGA POKOK PENJUALAN (HPP / COGS)", _SEC)
    sub("Project Based:")
    cp = [line(k, False) for k in ("cogs_p_direct", "cogs_p_nondirect", "cogs_p_material", "cogs_p_other")]
    sub("SaaS — %s:" % rep["title"])
    cs = [line(k, False) for k in ("cogs_s_direct", "cogs_s_cloud", "cogs_s_api", "cogs_s_material", "cogs_s_other")]
    r_cogs = total("TOTAL HPP", lambda L: "SUM(%s%d:%s%d)+SUM(%s%d:%s%d)" % (L, cp[0], L, cp[-1], L, cs[0], L, cs[-1]),
                   is_revenue=False)
    row[0] += 1
    r_gp = total("LABA KOTOR (GROSS PROFIT)", lambda L: "%s%d-%s%d" % (L, r_rev, L, r_cogs))
    row[0] += 1
    section("BIAYA OPERASIONAL (OPEX)", _SEC)
    ox = [line(k, False) for k in ("opex_salary", "opex_software", "opex_rent", "opex_util", "opex_event",
                                   "opex_meals", "opex_ent", "opex_medical", "opex_equipment", "opex_other")]
    r_opex = total("TOTAL OPEX", lambda L: "SUM(%s%d:%s%d)" % (L, ox[0], L, ox[-1]), is_revenue=False)
    row[0] += 1
    r_ebitda = total("EBITDA", lambda L: "%s%d-%s%d" % (L, r_gp, L, r_opex))
    r_da = line("da", False)
    r_ebit = total("EBIT", lambda L: "%s%d-%s%d" % (L, r_ebitda, L, r_da))
    r_tax = row[0]
    _box(ws, "B%d" % r_tax, "    Pajak Penghasilan (est. %s%%)" % ("%g" % (rep["pnl"]["tax_rate"] * 100)), None, color=_DEEP)
    # tax is estimated on the year-to-date and full-year profit, not month by month
    for k in (0, 1, 3):
        L = get_column_letter(ycol + k)
        ws["%s%d" % (L, r_tax)] = "=MAX(0,%s%d)*%s" % (L, r_ebit, rep["pnl"]["tax_rate"])
        ws["%s%d" % (L, r_tax)].number_format = NUM_FMT
    row[0] += 1
    total("LABA BERSIH (NET PROFIT)", lambda L: "%s%d-N(%s%d)" % (L, r_ebit, L, r_tax))

    # ================================================================ CASH FLOW
    ws = wb.create_sheet("CASH FLOW")
    cash = rep["cash"]
    rows = [r for r in cash["rows"] if r["closed"]]
    _box(ws, "A1:I1", "LAPORAN ARUS KAS — REAL-TIME TRACKER  |  %s YTD %d" % (rep["title"], rep["year"]),
         _NAVY, bold=True, size=13)
    _box(ws, "A2:I2", "Dari jurnal terposting pada proyek SBU, sampai %s. Saldo awal: %s."
         % (rep["through_label"] or "-", "{:,.0f}".format(cash["opening"]).replace(",", ".")), None, color=_GREY)
    first, lastr = 8, 7 + len(rows)
    _box(ws, "A4:A5", "Cash In Total", _GREEN, bold=True, align="center", wrap=True)
    _box(ws, "B4:C5", "=SUM(G%d:G%d)" % (first, max(first, lastr)), None, color=_DEEP, bold=True, fmt="#,##0")
    _box(ws, "D4:E5", "Cash Out Total", _RED, bold=True, align="center", wrap=True)
    _box(ws, "F4:G5", "=SUM(H%d:H%d)" % (first, max(first, lastr)), None, color=_DEEP, bold=True, fmt="#,##0")
    _box(ws, "H4:H5", "Net Cash", _NAVY, bold=True, align="center")
    _box(ws, "I4:I5", "=B4+F4", None, color=_DEEP, bold=True, fmt="#,##0")
    for i, (h, w) in enumerate((("Date", 12), ("Reference", 16), ("Description", 48), ("Expense", 26),
                                ("Category", 34), ("Tipe", 10), ("Cash In (IDR)", 16), ("Cash Out (IDR)", 16),
                                ("Running Balance (IDR)", 20)), start=1):
        _box(ws, "%s7" % get_column_letter(i), h, _NAVY, bold=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    for n, r in enumerate(rows):
        rr = first + n
        _date_cell(ws, rr, 1, r["date"])
        ws.cell(row=rr, column=2, value=r["entry_no"])
        ws.cell(row=rr, column=3, value=r["description"])
        ws.cell(row=rr, column=4, value=r["line"])
        ws.cell(row=rr, column=5, value=r["category"])
        ws.cell(row=rr, column=6, value=r["type"])
        _num(ws, rr, 7, r["cash_in"])
        _num(ws, rr, 8, r["cash_out"])
        ws.cell(row=rr, column=9, value=("=%s+G%d+H%d" % (cash["opening"], rr, rr)) if n == 0
                else "=I%d+G%d+H%d" % (rr - 1, rr, rr)).number_format = NUM_FMT
    ws.freeze_panes = "A8"
    r = lastr + 2
    _box(ws, "A%d:I%d" % (r, r), "Monthly Summary", _NAVY, bold=True)
    for i, h in enumerate(("Month", "Cash In Total", "Cash Out Total", "Net Cash Flow", "Ending Balance"), start=1):
        _box(ws, "%s%d" % (get_column_letter(i), r + 1), h, _GREY, bold=True, size=8)
    for n, s in enumerate(cash["months"][:T] if T else []):
        rr = r + 2 + n
        ws.cell(row=rr, column=1, value="%s %s" % (months[n], s["month"][:4]))
        _num(ws, rr, 2, s["in"])
        _num(ws, rr, 3, -s["out"])
        _num(ws, rr, 4, s["net"])
        _num(ws, rr, 5, s["ending"])

    # =========================================================== TRACKER PROYEK
    ws = wb.create_sheet("TRACKER PROYEK")
    tr = rep["tracker"]
    _box(ws, "A1:S1", "TRACKER PROYEK — BUDGET VS REALISASI  |  %s %d" % (rep["title"], rep["year"]),
         _TEAL, bold=True, size=13)
    _box(ws, "A2:S2", "Semua nilai dalam IDR. Realisasi dari jurnal terposting sampai %s; invoice dari Money Tracker "
         "atau isian SBU." % (rep["through_label"] or "-"), None, color=_GREY)
    for ref, text, fill in (("A3:B3", "PROYEK", _TEAL), ("C3:D3", "TIMELINE", _SEC), ("E3", "NILAI KONTRAK", _SEC),
                            ("F3:I3", "BUDGET BIAYA", _SEC), ("J3:M3", "REALISASI BIAYA", _SEC),
                            ("N3", "GROSS MARGIN", _SEC), ("O3:P3", "INVOICING", _SEC), ("Q3", "STATUS", _SEC),
                            ("R3", "RISIKO", _SEC), ("S3", "NOTE", _SEC)):
        _box(ws, ref, text, fill, bold=True, size=8, align="center")
    heads = ["Nama Proyek", "Klien", "Mulai", "Selesai", "Nilai Kontrak", "Bud. Direct", "Bud. Non Direct",
             "Bud. Materials", "Bud. Fixed/Misc", "Act. Direct", "Act. Non Direct", "Act. Material",
             "Act. Fixed/Misc", "Gross Margin", "Invoice", "Terbayar", "Status", "Risiko", "Note"]
    widths = [34, 26, 11, 16, 16, 15, 15, 15, 15, 15, 15, 15, 15, 16, 16, 16, 14, 16, 60]
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        _box(ws, "%s4" % get_column_letter(i), h, _GREY, bold=True, size=8, align="center", wrap=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    bk = ("direct", "nondirect", "material", "misc")
    for n, t in enumerate(tr["rows"]):
        rr = 5 + n
        ws.cell(row=rr, column=1, value="%s — %s" % (t["code"], t["name"])).font = BOLD
        ws.cell(row=rr, column=2, value=t["client"])
        _date_cell(ws, rr, 3, t["start"])
        _date_cell(ws, rr, 4, t["end"])
        _num(ws, rr, 5, t["contract_value"])
        for k, key in enumerate(bk):
            _num(ws, rr, 6 + k, t["budget"][key])
            _num(ws, rr, 10 + k, t["actual"][key])
        ws.cell(row=rr, column=14, value='=IF(E%d>0,(E%d-SUM(J%d:M%d))/E%d,"")' % (rr, rr, rr, rr, rr)).number_format = "0.00%"
        _num(ws, rr, 15, t["invoiced"])
        _num(ws, rr, 16, t["paid"])
        ws.cell(row=rr, column=17, value=t["status"])
        ws.cell(row=rr, column=18, value=t["risk"])
        ws.cell(row=rr, column=19, value=t["note"]).alignment = Alignment(wrap_text=True, vertical="top")
    lr = 4 + max(1, len(tr["rows"]))
    r = lr + 3
    _box(ws, "A%d:R%d" % (r, r), "RINGKASAN PIPELINE & COLLECTION", _NAVY, bold=True)
    for ref, label, formula in (("A%d:C%d", "Total Nilai Pipeline", "=SUM(E5:E%d)" % lr),
                                ("E%d:H%d", "Total Invoice Terkirim", "=SUM(O5:O%d)" % lr),
                                ("J%d:M%d", "Total Terbayar", "=SUM(P5:P%d)" % lr),
                                ("O%d:Q%d", "Total Outstanding", None)):
        _box(ws, ref % (r + 1, r + 1), label, "F0F4F8", color=_DEEP, bold=True)
    ws["D%d" % (r + 1)] = "=SUM(E5:E%d)" % lr
    ws["I%d" % (r + 1)] = "=SUM(O5:O%d)" % lr
    ws["N%d" % (r + 1)] = "=SUM(P5:P%d)" % lr
    ws["R%d" % (r + 1)] = "=I%d-N%d" % (r + 1, r + 1)
    for c in ("D", "I", "N", "R"):
        ws["%s%d" % (c, r + 1)].number_format = "#,##0"
        ws["%s%d" % (c, r + 1)].font = BOLD
    ws.freeze_panes = "B5"

    # ====================================================== INDIKATOR KESEHATAN
    ws = wb.create_sheet("INDIKATOR KESEHATAN")
    _box(ws, "A1:F1", "INDIKATOR KESEHATAN KEUANGAN — SISTEM PERINGATAN  |  %s %d" % (rep["title"], rep["year"]),
         _NAVY, bold=True, size=13)
    _box(ws, "A2:F2", "Nilai dihitung dari buku besar per %s; nilai yang diisi manual ditandai. "
         "Perbarui setiap tutup bulan." % (rep["through_label"] or "-"), None, color=_GREY)
    _box(ws, "A4", "LEGENDA STATUS:", "F0F4F8", color=_DEEP, bold=True)
    for ref, key in (("B4", "safe"), ("C4", "watch"), ("D4", "critical")):
        st, bg, fg = _STATUS_ID[key]
        _box(ws, ref, st, bg, color=fg, bold=True)
    for i, (h, w) in enumerate((("Indikator Kesehatan", 30), ("Nilai Aktual", 20), ("Batas AMAN (🟢)", 20),
                                ("Batas PERHATIAN (🟡)", 22), ("Status", 16), ("Interpretasi & Tindakan", 80)), start=1):
        _box(ws, "%s6" % get_column_letter(i), h, _NAVY, bold=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    r = 7
    group_label = {"liquidity": "LIKUIDITAS & SOLVABILITAS", "profit": "PROFITABILITAS",
                   "saas": "PERFORMA SaaS — %s" % title, "ops": "OPERASIONAL & PROYEK"}
    for g in rep["groups"]:
        _box(ws, "A%d:F%d" % (r, r), group_label.get(g["key"], g["label"]), _NAVY if g["key"] == "liquidity" else _SEC, bold=True)
        r += 1
        for ind in [i for i in rep["indicators"] if i["group"] == g["key"]]:
            st, bg, fg = _STATUS_ID[ind["status"]]
            _box(ws, "A%d" % r, ind["label"] + (" (manual)" if ind["source"] == "entered" else ""), "F0F4F8", color=_DEEP, bold=True)
            _box(ws, "B%d" % r, ind["value"] if ind["value"] is not None else "-", bg, color=fg,
                 fmt=_UNIT_FMT.get(ind["unit"]) if ind["value"] is not None else None)
            _box(ws, "C%d" % r, _limit_text(ind["safe"], ind["unit"], ind["direction"]), None, color=_GREY)
            _box(ws, "D%d" % r, _limit_text(ind["attention"], ind["unit"], ind["direction"]), None, color=_GREY)
            _box(ws, "E%d" % r, st, bg, color=fg, bold=True)
            _box(ws, "F%d" % r, ind["note"], None, color=_GREY, wrap=True)
            ws.row_dimensions[r].height = 28
            r += 1
        r += 1
    r += 1
    _box(ws, "A%d:F%d" % (r, r), "⚡  PRIORITAS TINDAKAN DIREKSI — %s" % rep["quarter"], _RED, bold=True, size=10)
    r += 1
    _box(ws, "A%d" % r, "Urgensi", "F0F4F8", color=_DEEP, bold=True)
    _box(ws, "B%d:F%d" % (r, r), "Tindakan yang Diperlukan", "F0F4F8", color=_DEEP, bold=True)
    labels = {u["key"]: u["label"] for u in rep["urgencies"]}
    style = {"now": ("FADBD8", _RED), "week": ("FEF9E7", _AMBER), "month": ("FEF9E7", _AMBER),
             "quarter": ("D6EAF8", _SEC)}
    acts = rep["actions"] or rep["suggested_actions"]
    for a in acts:
        r += 1
        bg, fg = style.get(a["urgency"], style["now"])
        _box(ws, "A%d" % r, labels.get(a["urgency"], a["urgency"]), bg, color=fg, bold=True)
        _box(ws, "B%d:F%d" % (r, r), a["text"], None, color=_DEEP, wrap=True)
        ws.row_dimensions[r].height = 30
    if not rep["actions"] and acts:
        r += 1
        _box(ws, "A%d:F%d" % (r, r), "Saran otomatis dari indikator KRITIS — isi tindakan direksi di SBU → Report settings.",
             None, color=_GREY, size=8)

    # ============================================================= KINERJA SaaS
    ws = wb.create_sheet("KINERJA SaaS")
    sa = rep["saas"]
    tl = sa["tiles"]
    _box(ws, "A1:M1", "KINERJA SaaS — %s  |  Subscription Metrics %d" % (rep["title"], rep["year"]), _BLUE,
         bold=True, size=13)
    _box(ws, "A2:M2", "Pelanggan dan MRR aktual diisi di SBU → Report settings; kosong = diambil dari pendapatan "
         "proyek SaaS di buku besar.", None, color=_GREY)
    ws.column_dimensions["A"].width = 32
    for c in range(2, 14):
        ws.column_dimensions[get_column_letter(c)].width = 13
    for ref, label, vref, value, fmt in (("A4:C4", "ARR (Annualized)", "A5:C5", tl["arr"], "#,##0"),
                                          ("D4:F4", "MRR Terkini (%s)" % (months[T - 1] if T else "-"), "D5:F5", tl["mrr"], '"Rp"#,##0'),
                                          ("G4:I4", "Total Customer", "G5:I5", "%s Pelanggan" % tl["customers"], None),
                                          ("J4:K4", "Churn Rate %s" % tl["quarter"], "J5:K5", tl["churn_quarter"], "0.0%"),
                                          ("L4:M4", "ARPU Rata-rata", "L5:M5", tl["arpu"], '"Rp"#,##0')):
        _box(ws, ref, label, _BLUE if ref[0] in "AGL" else _CYAN, color="BDC3C7", size=8)
        _box(ws, vref, value if value is not None else "-", _BLUE if ref[0] in "AGL" else _CYAN, bold=True,
             size=14, fmt=fmt if isinstance(value, (int, float)) else None, align="center")
    _box(ws, "A7", "METRIK", _BLUE, bold=True)
    for m in range(12):
        _box(ws, "%s7" % get_column_letter(2 + m), months[m], _BLUE, bold=True, align="center")
    r = 8

    def block(label):
        nonlocal r
        _box(ws, "A%d:M%d" % (r, r), label, _BLUE, bold=True)
        r += 1

    def mrow(label, key, fmt="#,##0", upto_closed=False):
        nonlocal r
        _box(ws, "A%d" % r, label, "F0F4F8", color=_DEEP)
        for m, s in enumerate(sa["rows"]):
            v = s.get(key)
            if v is None or (upto_closed and not s["closed"]):
                continue
            c = ws.cell(row=r, column=2 + m, value=v)
            c.number_format = fmt
        r += 1

    block("CUSTOMER (PELANGGAN)")
    mrow("Customer Awal Bulan", "start", "0")
    mrow("Pelanggan Baru (Akuisisi)", "new", "0")
    mrow("Churned (Berhenti)", "churned", "0")
    mrow("Customer Akhir Bulan", "end", "0")
    r += 1
    block("REVENUE (MRR / ARR)")
    mrow("MRR Budget (Rp)", "mrr_budget")
    mrow("MRR Aktual (Rp)", "mrr_actual", upto_closed=True)
    mrow("ARR Annualized (Rp)", "arr", upto_closed=True)
    mrow("Variance MRR (Rp)", "variance", upto_closed=True)
    mrow("MRR Achievement %", "achievement", "0.0%", upto_closed=True)
    r += 1
    block("CHURN, GROWTH & UNIT ECONOMICS")
    mrow("Churn Rate (%)", "churn", "0.0%")
    mrow("MRR Growth MoM (%)", "growth", "0.0%", upto_closed=True)
    mrow("ARPU (Avg Rev / User, Rp)", "arpu", upto_closed=True)
    mrow("LTV est. (asumsi churn %s%%, Rp)" % ("%g" % (sa["ltv_churn"] * 100)), "ltv", upto_closed=True)
    mrow("CAC (Biaya Akuisisi, Rp)", "cac")
    mrow("LTV/CAC Ratio", "ltv_cac", "0.0")
    _box(ws, "A%d:M%d" % (r, r), "💡  Benchmark SaaS Sehat: Churn < 5%/bln | LTV/CAC > 3x | NRR > 100% | "
         "MRR Growth > 15%/bln.", "FEF9E7", color=_AMBER, size=9)
    return _to_bytes(wb)
