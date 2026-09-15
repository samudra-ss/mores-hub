"""SBU model workbook: export, edit in Excel terms, import.

The round trip must be lossless, an edit must land exactly, and a bad file must
change NOTHING. Runs against a fresh demo database in a temporary folder.

    python apps/erp/tests/test_sbu_excel.py
"""
import datetime
import io
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
TMP = tempfile.mkdtemp(prefix="sbu_xl_")
os.environ["MORES_HV_DATA_DIR"] = TMP

import database  # noqa: E402
database.create_database(database.DEFAULT_DB, seed_demo=True)
import excel_io as X  # noqa: E402
import server  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

ok = fail = 0


def check(name, cond, detail=""):
    global ok, fail
    if cond:
        ok += 1
        print("  PASS  %-10s %s" % (name, detail))
    else:
        fail += 1
        print("  FAIL  %-10s %s" % (name, detail))


def raises(fn, *args):
    try:
        fn(*args)
        return False
    except ValueError:
        return True


print("=== reading cells the way people fill them in ===")
check("MONTH", [X._sbu_month_in(v) for v in ("09/2026", "9/2026", datetime.datetime(2026, 9, 1), "2026-09",
                                             "01/09/2026", None, "")] == ["2026-09"] * 5 + [None, None],
      "MM/YYYY, a real Excel date, YYYY-MM and DD/MM/YYYY all read as 2026-09")
check("MONTH-BAD", raises(X._sbu_month_in, "13/2026") and raises(X._sbu_month_in, "September"),
      "an impossible month is an error, not a guess")
wb = Workbook()
ws = wb.active
ws["A1"] = 0.02
ws["A1"].number_format = "0.00%"
ws["A2"] = 2
ws["A3"] = "2%"
check("PCT", [round(X._sbu_pct_in(ws[c]), 6) for c in ("A1", "A2", "A3")] == [0.02, 0.02, 0.02],
      "an Excel percentage, a plain 2 and the text 2% all mean 2%")
check("AMOUNT", [X._sbu_money_in(v) for v in (1200000, "1.200.000", "1,200,000", "Rp 1.200.000", "")]
      == [1200000.0] * 4 + [0.0], "numbers, Indonesian and English separators, and Rp")

c = server.app.test_client()
assert c.post("/api/login", json={"username": "admin", "password": "MoresMores2018"}).status_code == 200
pid = c.get("/api/products").get_json()[0]["id"]


def detail(p):
    d = c.get("/api/products/%d" % p).get_json()
    pick = lambda rows, keys: [tuple(r.get(k) for k in keys) for r in rows]
    return {
        "row": tuple(d.get(k) for k in ("name", "code", "company_id", "stage", "launch_month", "target_year",
                                         "actual_through", "horizon_end", "burn_budget", "notes")),
        "assumptions": {k: v for k, v in d["assumptions"].items() if k != "category_map"},
        "lines": pick(d["lines"], ("label", "category", "monthly_amount", "basis", "start_month", "end_month",
                                   "escalation_annual")),
        "oneoffs": pick(d["oneoffs"], ("month", "flow", "label", "amount")),
        "capex": pick(d["capex"], ("month", "label", "amount", "life_months")),
        "projects": sorted(d["project_ids"]),
    }


def upload(url, data, extra=None):
    form = {"file": (io.BytesIO(data), "sbu.xlsx")}
    form.update(extra or {})
    return c.post(url, data=form, content_type="multipart/form-data")


def key_row(ws, key):
    return next(r for r in range(5, ws.max_row + 1) if ws.cell(row=r, column=4).value == key)


def saved(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


print("=== export, then import the same file: nothing may change ===")
before = detail(pid)
an_before = c.get("/api/products/%d/analysis" % pid).get_json()
r = c.get("/api/products/%d/workbook" % pid)
check("EXPORT", r.status_code == 200 and r.data[:2] == b"PK", "the SBU downloads as an .xlsx")
xl = r.data
names = load_workbook(io.BytesIO(xl)).sheetnames
check("SHEETS", names[:6] == ["SBU", "Cost lines", "One-offs", "CAPEX", "Projects", "Project list"], str(names))
res = upload("/api/products/%d/import" % pid, xl).get_json()
check("IMPORT", res.get("ok") is True, str(res.get("summary") or res.get("errors")))
after = detail(pid)
for part in ("row", "assumptions", "lines", "oneoffs", "capex", "projects"):
    check("SAME-" + part[:5].upper(), before[part] == after[part],
          "%s identical after the round trip" % part if before[part] == after[part]
          else "before %r / after %r" % (before[part], after[part]))
an_after = c.get("/api/products/%d/analysis" % pid).get_json()
check("SAME-MODEL", [(m["revenue"], m["opex"], m["cum_cash"]) for m in an_before["months"]]
      == [(m["revenue"], m["opex"], m["cum_cash"]) for m in an_after["months"]],
      "every month of the analysis is identical")

print("=== an edit made in Excel lands exactly ===")
wb = load_workbook(io.BytesIO(xl))
sheet = wb["SBU"]
sheet.cell(row=key_row(sheet, "churn_monthly"), column=2).value = 0.05
sheet.cell(row=key_row(sheet, "name"), column=2).value = "Sentimind (edited in Excel)"
wl = wb["Cost lines"]
new_row = 5 + len(before["lines"])
for col, val in enumerate(["QA marketing push", "marketing", 5000000, "per month", "01/2027", "06/2027"], start=1):
    wl.cell(row=new_row, column=col).value = val
wl.cell(row=new_row, column=7).value = 0.1
res = upload("/api/products/%d/import" % pid, saved(wb)).get_json()
edited = detail(pid)
check("EDIT", res.get("ok") and edited["assumptions"]["churn_monthly"] == 0.05
      and edited["row"][0] == "Sentimind (edited in Excel)", "churn 5% and the new name are saved")
check("EDIT-LINE", len(edited["lines"]) == len(before["lines"]) + 1 and edited["lines"][-1]
      == ("QA marketing push", "marketing", 5000000.0, "fixed", "2027-01", "2027-06", 0.1),
      "the added cost line: %r" % (edited["lines"][-1],))

print("=== a bad file changes nothing ===")
wb = load_workbook(io.BytesIO(saved(wb)))
wb["Cost lines"].cell(row=5, column=2).value = "snacks"
wb["One-offs"].cell(row=5, column=1).value = "13/2026"
res = upload("/api/products/%d/import" % pid, saved(wb)).get_json()
errs = res.get("errors") or []
check("BAD-REJECT", res.get("ok") is False and any("Cost lines row 5" in e for e in errs)
      and any("One-offs row 5" in e for e in errs), "; ".join(errs)[:160])
check("BAD-UNTOUCHED", detail(pid) == edited, "the SBU is exactly as it was before the bad upload")
wb = load_workbook(io.BytesIO(xl))
sheet = wb["SBU"]
sheet.cell(row=key_row(sheet, "churn_monthly"), column=2).value = 1.5          # 150% a month
res = upload("/api/products/%d/import" % pid, saved(wb)).get_json()
check("BAD-RULE", res.get("ok") is False and any("Churn" in e for e in res.get("errors") or [])
      and detail(pid) == edited, "a value the settings tab would refuse is refused here too, nothing saved")
res = upload("/api/products/%d/import" % pid, b"this is not a workbook").get_json()
check("NOT-XLSX", res.get("ok") is False and "could not be opened" in (res.get("errors") or [""])[0], "")
empty = Workbook()
empty.active.title = "Budget"
res = upload("/api/products/%d/import" % pid, saved(empty)).get_json()
check("NO-SBU-SHEET", res.get("ok") is False and "no sheet named 'SBU'" in (res.get("errors") or [""])[0], "")

print("=== a new SBU straight from the blank template ===")
r = c.get("/api/products/template")
check("TEMPLATE", r.status_code == 200 and r.data[:2] == b"PK", "blank template downloads")
wb = load_workbook(io.BytesIO(r.data))
sheet = wb["SBU"]
for key, val in (("name", "SBU from Excel"), ("company", "MDA"), ("launch_month", "03/2027"),
                 ("target_year", 2028), ("arpu", 150000)):
    sheet.cell(row=key_row(sheet, key), column=2).value = val
churn = sheet.cell(row=key_row(sheet, "churn_monthly"), column=2)
churn.value = 3
churn.number_format = "General"                     # typed as a plain 3
wb["Projects"].cell(row=5, column=1).value = "PRJ-NX"
count = len(c.get("/api/products").get_json())
r = upload("/api/products/import", saved(wb), {"company_id": "1"})
res = r.get_json()
check("NEW", r.status_code == 201 and res.get("ok"), str(res.get("summary") or res.get("errors")))
new = detail(res["id"]) if res.get("ok") else None
nx = next((p for p in c.get("/api/products/%d" % pid).get_json()["projects"] if p["code"] == "PRJ-NX"), None)
check("NEW-FIELDS", new and new["row"][0] == "SBU from Excel" and new["row"][4] == "2027-03" and new["row"][5] == 2028
      and new["assumptions"]["arpu"] == 150000 and new["assumptions"]["churn_monthly"] == 0.03,
      "name, launch 03/2027, target 2028, ARPU and a plain-typed 3%% churn")
check("NEW-COMPANY", new and new["row"][2] == next(
    co["id"] for co in c.get("/api/companies").get_json() if co["code"] == "MDA"), "company code MDA resolved")
check("NEW-PROJECT", new and nx and new["projects"] == [nx["id"]], "project PRJ-NX linked by its code")
check("NEW-LIST", len(c.get("/api/products").get_json()) == count + 1, "one more SBU in the list")

wb["Projects"].cell(row=5, column=1).value = "NO-SUCH-PROJECT"
res = upload("/api/products/import", saved(wb), {"company_id": "1"}).get_json()
check("NEW-BADPRJ", res.get("ok") is False and any("NO-SUCH-PROJECT" in e for e in res.get("errors") or [])
      and len(c.get("/api/products").get_json()) == count + 1, "an unknown project code creates nothing")

print()
print("RESULT: %d passed, %d failed" % (ok, fail))
shutil.rmtree(TMP, ignore_errors=True)
sys.exit(1 if fail else 0)
