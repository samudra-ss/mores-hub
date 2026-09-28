"""Per-project budgets through Excel: template, import, export, round trip.

The rules that must not quietly change: an import aimed at a project puts its
rows THERE (a blank Project Code is not company-level money), a row naming a
different project is refused rather than written somewhere nobody looked, and a
re-imported month keeps the week shape finance set.

    python apps/erp/tests/test_budget_project_excel.py
"""
import io
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
TMP = tempfile.mkdtemp(prefix="budget_xl_")
os.environ["MORES_HV_DATA_DIR"] = TMP

import database  # noqa: E402
database.create_database(database.DEFAULT_DB, seed_demo=True)
import server  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

YEAR = 2026
ok = fail = 0


def check(name, cond, detail=""):
    global ok, fail
    if cond:
        ok += 1
        print("  PASS  %-12s %s" % (name, detail))
    else:
        fail += 1
        print("  FAIL  %-12s %s" % (name, detail))


c = server.app.test_client()
assert c.post("/api/login", json={"username": "admin", "password": "MoresMores2018"}).status_code == 200
company = next(x for x in c.get("/api/companies").get_json() if x["code"] == "MDA")
cid = company["id"]
projects = [p for p in c.get("/api/projects?company_id=%d" % cid).get_json()]
A, B = projects[0], projects[1]
ACC = "6100"


def grid():
    return c.get("/api/budgets?company_id=%d&year=%d" % (cid, YEAR)).get_json()["rows"]


def row_for(project_id, code=ACC):
    return next((r for r in grid() if r["code"] == code and r["project_id"] == project_id), None)


def book(rows):
    """A budget workbook the way a person fills the template in."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Budget Import Template"])
    ws.append([])
    ws.append([])
    ws.append(["Account Code", "Project Code (optional)"] + ["M%d" % m for m in range(1, 13)])
    for code, pcode, amounts in rows:
        ws.append([code, pcode] + list(amounts))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def imp(data, **fields):
    form = {"file": (io.BytesIO(data), "budget.xlsx"), "company_id": str(cid), "year": str(YEAR)}
    form.update({k: str(v) for k, v in fields.items()})
    return c.post("/api/import/budget", data=form, content_type="multipart/form-data").get_json()


print("=== the template carries the project codes ===")
r = c.get("/api/templates/budget?year=%d&company_id=%d&project_id=%d" % (YEAR, cid, A["id"]))
wb = load_workbook(io.BytesIO(r.data))
check("TPL-SHEETS", r.status_code == 200 and "Project list" in wb.sheetnames, str(wb.sheetnames))
listed = {wb["Project list"].cell(row=i, column=1).value for i in range(4, 4 + len(projects) + 2)}
check("TPL-CODES", {p["code"] for p in projects} <= listed, "every project code is listed to copy from")
check("TPL-PREFILL", wb["Budget"].cell(row=5, column=2).value == A["code"],
      "the open project is already in column B")
plain = load_workbook(io.BytesIO(c.get("/api/templates/budget?year=%d" % YEAR).data))
check("TPL-PLAIN", "Project list" not in plain.sheetnames, "without a company it stays the plain template")

print("=== importing INTO a project ===")
# the demo company already budgets at company level, so the point is that the
# import leaves that alone rather than that nothing is there
comp_before = sum((row_for(None) or {"amounts": [0]})["amounts"])
res = imp(book([(ACC, "", [1000000] * 12)]), project_id=A["id"])
ra = row_for(A["id"])
check("INTO-PRJ", res["saved_rows"] == 1 and ra and sum(ra["amounts"]) == 12000000,
      "a blank Project Code lands on the project the import was for, not at company level")
check("INTO-NONE", sum((row_for(None) or {"amounts": [0]})["amounts"]) == comp_before,
      "and the company-level line for the same account is untouched (%s)" % comp_before)
check("INTO-SAYS", res["by_project"] == {A["code"]: 1}, "the result says where the rows went: %s" % res["by_project"])

res = imp(book([(ACC, B["code"], [9999] * 12)]), project_id=A["id"])
check("WRONG-PRJ", res["saved_rows"] == 0 and res["errors"] and B["code"] in res["errors"][0]
      and row_for(B["id"]) is None,
      "a row naming another project is refused, not written: %s" % (res["errors"] or [""])[0][:74])

res = imp(book([(ACC, "", [2000000] * 12)]))
check("NO-PRJ", res["saved_rows"] == 1 and row_for(None) and sum(row_for(None)["amounts"]) == 24000000,
      "with no project chosen, a blank code is still a company-level line")
res = imp(book([(ACC, B["code"], [3000000] * 12)]))
check("BY-CODE", res["saved_rows"] == 1 and sum(row_for(B["id"])["amounts"]) == 36000000,
      "and a code in the file puts the row on that project")
res = imp(book([("9999", "", [1] * 12), (ACC, "NO-SUCH", [1] * 12)]))
check("BAD-ROWS", res["saved_rows"] == 0 and len(res["errors"]) == 2
      and all("Row 5" in e or "Row 6" in e for e in res["errors"]),
      "an unknown account or project is reported with its row number")

print("=== exporting every project at once ===")
r = c.get("/api/export/budget?company_id=%d&year=%d&project_id=all" % (cid, YEAR))
check("EXP-ALL-OK", r.status_code == 200 and "all_projects" in r.headers.get("Content-Disposition", ""),
      r.headers.get("Content-Disposition", ""))
ws = load_workbook(io.BytesIO(r.data)).active
body = [[ws.cell(row=i, column=j).value for j in range(1, 17)]
        for i in range(5, ws.max_row + 1) if ws.cell(row=i, column=1).value]
prj_rows = [x for x in grid() if x["project_id"]]
check("EXP-ALL-ROWS", len(body) == len(prj_rows) and {b[1] for b in body} >= {A["code"], B["code"]},
      "every project row in one file (%d), each carrying its project code" % len(body))
check("EXP-ALL-NONE", all(b[1] for b in body), "and no company-level row sneaks in")
one = load_workbook(io.BytesIO(c.get(
    "/api/export/budget?company_id=%d&year=%d&project_id=%d" % (cid, YEAR, A["id"])).data)).active
mine = [x for x in prj_rows if x["project_id"] == A["id"]]
check("EXP-ONE", one["A1"].value.endswith(A["code"]) and one.max_row == 4 + len(mine),
      "one project exports on its own: %r, %d rows" % (one["A1"].value, len(mine)))
comp = load_workbook(io.BytesIO(c.get(
    "/api/export/budget?company_id=%d&year=%d" % (cid, YEAR)).data)).active
check("EXP-COMPANY", all(comp.cell(row=i, column=2).value in (None, "")
                         for i in range(5, comp.max_row + 1)),
      "and the company-level export carries no project rows")

print("=== round trip keeps the week shape ===")
conn = database.get_db(database.DEFAULT_DB)
acc_id = conn.execute("SELECT id FROM accounts WHERE company_id=? AND code=?", (cid, ACC)).fetchone()[0]
for wk, amt in ((1, 0), (4, 4000000)):                    # March is deliberately W4-heavy
    database.upsert_budget(conn, cid, acc_id, A["id"], YEAR, 3, amt, week=wk)
conn.commit()
before = row_for(A["id"])["weeks"][2]
data = c.get("/api/export/budget?company_id=%d&year=%d&project_id=all" % (cid, YEAR)).data
res = imp(data)
after = row_for(A["id"])
check("RT-SHAPE", before == [0, 0, 0, 4000000] and after["weeks"][2] == before,
      "a month sitting in W4 comes back in W4, not flattened into W1: %s" % after["weeks"][2])
check("RT-TOTAL", sum(after["amounts"]) == 4000000 + 11000000 and res["errors"] == [],
      "and the yearly total is unchanged (%s)" % sum(after["amounts"]))
expect = {}
for x in prj_rows:
    code = A["code"] if x["project_id"] == A["id"] else B["code"]
    expect[code] = expect.get(code, 0) + 1
check("RT-SPLIT", res["by_project"] == expect,
      "the whole per-project set round-trips through one workbook: %s" % res["by_project"])
conn.close()

print()
print("RESULT: %d passed, %d failed" % (ok, fail))
shutil.rmtree(TMP, ignore_errors=True)
sys.exit(1 if fail else 0)
