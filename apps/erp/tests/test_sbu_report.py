"""v1.10: the SBU finance report (the NX-Sentimind workbook, from the ledger).

Covers the rules that must not quietly change: every rupiah on the SBU's
projects lands on exactly one P&L line, the project KIND decides revenue and
HPP, settings the ledger cannot supply survive an Excel model import, and the
exported workbook has the template's six sheets.

    python apps/erp/tests/test_sbu_report.py
"""
import io
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
TMP = tempfile.mkdtemp(prefix="sbu_rep_")
os.environ["MORES_HV_DATA_DIR"] = TMP

import database  # noqa: E402
database.create_database(database.DEFAULT_DB, seed_demo=True)
import sbu_report as R  # noqa: E402
import server  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ok = fail = 0


def check(name, cond, detail=""):
    global ok, fail
    if cond:
        ok += 1
        print("  PASS  %-12s %s" % (name, detail))
    else:
        fail += 1
        print("  FAIL  %-12s %s" % (name, detail))


c = server.app.test_client()
assert c.post("/api/login", json={"username": "admin", "password": "MoresMores2018"}).status_code == 200
prods = c.get("/api/products").get_json()
pid = prods[0]["id"]
projects = {p["code"]: p for p in c.get("/api/projects?company_id=all").get_json()}
nx, lic = projects["PRJ-NX"], projects["PRJ-LIC"]
assert c.put("/api/products/%d" % pid, json={"project_ids": [nx["id"], lic["id"]]}).status_code == 200
Y = 2026


def report(**settings):
    if settings:
        r = c.put("/api/products/%d/report" % pid, json=settings)
        assert r.status_code == 200, r.get_json()
    return c.get("/api/products/%d/report?year=%d" % (pid, Y)).get_json()


def ledger(kind, through, project_ids):
    """Straight from the tables, independent of the report code."""
    conn = database.get_db(database.DEFAULT_DB)
    sign = "jl.credit - jl.debit" if kind == "revenue" else "jl.debit - jl.credit"
    v = conn.execute(
        "SELECT COALESCE(SUM(%s), 0) FROM journal_lines jl JOIN journal_entries je ON je.id = jl.entry_id"
        " JOIN accounts a ON a.id = jl.account_id WHERE je.status = 'posted' AND a.type = ?"
        " AND jl.project_id IN (%s) AND je.date BETWEEN ? AND ?"
        % (sign, ",".join("?" * len(project_ids))),
        [kind] + list(project_ids) + ["%d-01-01" % Y, "%d-%02d-31" % (Y, through)]).fetchone()[0]
    conn.close()
    return round(v, 2)


print("=== every rupiah lands on exactly one line ===")
rep = report(through="%d-06" % Y)
T = rep["through"]
check("THROUGH", T == 6 and rep["quarter"] == "Q2 %d" % Y, "closed through %s, %s" % (rep["through_label"], rep["quarter"]))
ytd = rep["pnl"]["ytd"]["actual"]
want_rev = ledger("revenue", 6, [nx["id"], lic["id"]])
want_exp = ledger("expense", 6, [nx["id"], lic["id"]])
check("REVENUE", abs(ytd["revenue"] - want_rev) < 0.5, "report %s = ledger %s" % (ytd["revenue"], want_rev))
check("ALL-COST", abs(ytd["cogs"] + ytd["opex"] + ytd["da"] - want_exp) < 0.5,
      "HPP + OPEX + D&A %s = every cost line %s" % (round(ytd["cogs"] + ytd["opex"] + ytd["da"], 2), want_exp))
lines_sum = sum(l["actual_ytd"] for l in rep["pnl"]["lines"])
check("LINES", abs(lines_sum - (want_rev + want_exp)) < 1, "the lines add back to the ledger")
check("NET", abs(ytd["net"] - (ytd["ebit"] - max(0, ytd["ebit"]) * 0.22)) < 0.5, "net = EBIT less 22% tax on a profit")

print("=== the project kind decides revenue and HPP ===")
check("DEFAULT", ytd["rev_saas"] == 0 and ytd["rev_project"] == ytd["revenue"], "every project starts as a client project")
nx_rev = ledger("revenue", 6, [nx["id"]])
rep2 = report(through="%d-06" % Y, projects={str(nx["id"]): {"kind": "saas"}})
y2 = rep2["pnl"]["ytd"]["actual"]
check("KIND", abs(y2["rev_saas"] - nx_rev) < 0.5 and abs(y2["revenue"] - ytd["revenue"]) < 0.5,
      "PRJ-NX's %s moves to SaaS; the total does not change" % nx_rev)
check("KIND-HPP", y2["cogs_saas"] > 0 and abs(y2["cogs"] - ytd["cogs"]) < 0.5, "its HPP moves with it")

print("=== account-code mapping ===")
before = {l["key"]: l["actual_ytd"] for l in rep2["pnl"]["lines"]}
rep3 = report(through="%d-06" % Y, projects={str(nx["id"]): {"kind": "saas"}},
              lines={"cogs_p_material": {"codes": "-"}, "cogs_s_material": {"codes": "-"}})
after = {l["key"]: l["actual_ytd"] for l in rep3["pnl"]["lines"]}
check("REMAP", after["cogs_p_material"] == 0 and abs(after["cogs_p_other"] - before["cogs_p_other"]
      - before["cogs_p_material"]) < 0.5, "a line emptied with '-' passes its accounts to Others Cost")
check("REMAP-SUM", abs(rep3["pnl"]["ytd"]["actual"]["cogs"] - y2["cogs"]) < 0.5, "and HPP still totals the same")
lines = R.resolved_lines(R.load_settings({}))
check("CATCH-ALL", R.line_for("6900", "expense", "project", lines) == "opex_other"
      and R.line_for("6500", "expense", "project", lines) == "da"
      and R.line_for("5200", "expense", "saas", lines) == "cogs_s_other", "6900 -> other opex, 6500 -> D&A, 5200 -> SaaS other")

print("=== indicators ===")
rep4 = report(through="%d-06" % Y, projects={str(nx["id"]): {"kind": "saas"}},
              indicators={"gross_margin": {"safe": 0.99, "attention": 0.98}, "nrr": {"actual": 1.04, "note": "by hand"}})
ind = {i["key"]: i for i in rep4["indicators"]}
check("LIMITS", ind["gross_margin"]["status"] == "critical" and ind["gross_margin"]["safe"] == 0.99,
      "a stricter target turns gross margin %s KRITIS" % ind["gross_margin"]["value"])
check("OVERRIDE", ind["nrr"]["value"] == 1.04 and ind["nrr"]["source"] == "entered" and ind["nrr"]["status"] == "safe",
      "NRR typed by hand is used and marked")
check("TWELVE", len(rep4["dashboard_indicators"]) == 12 and len(rep4["tiles"]) == 8, "12 dashboard indicators, 8 tiles")

print("=== project tracker ===")
rep5 = report(through="%d-06" % Y, projects={str(lic["id"]): {"kind": "project", "contract_value": 3000000000,
                                                                   "invoiced": 2000000000, "risk": "Tinggi"}})
row = next(t for t in rep5["tracker"]["rows"] if t["project_id"] == lic["id"])
check("TRACK", row["contract_value"] == 3000000000 and row["invoiced_source"] == "entered"
      and abs(row["outstanding"] - max(0, 2000000000 - row["paid"])) < 0.5,
      "contract and invoice typed in settings; outstanding %s" % row["outstanding"])
check("TRACK-COST", abs(sum(row["actual"].values()) - ledger("expense", 6, [lic["id"]])) < 0.5,
      "the four cost columns hold every cost line on the project")
check("GM", abs(row["gross_margin"] - (3000000000 - row["actual_total"]) / 3000000000) < 0.0001,
      "gross margin on the contract value")

print("=== SaaS performance ===")
rep6 = report(through="%d-06" % Y, opening_customers=5,
              saas_months={"%d-01" % Y: {"new": 3, "churned": 1, "mrr_actual": 9000000, "cac": 6000000},
                           "%d-02" % Y: {"new": 2, "churned": 0, "mrr_actual": 12000000}})
s = rep6["saas"]["rows"]
check("ROLL", s[0]["start"] == 5 and s[0]["end"] == 7 and s[1]["start"] == 7 and s[1]["end"] == 9,
      "customers carry from one month to the next")
check("CHURN", abs(s[0]["churn"] - 0.2) < 1e-9 and abs(s[1]["growth"] - 1 / 3) < 1e-3,
      "churn 1 of 5, growth 9 -> 12 jt")
check("CAC", s[0]["cac"] == 2000000, "6 jt spent / 3 new customers")

print("=== validation and roles ===")
bad = c.put("/api/products/%d/report" % pid, json={"through": "June"})
check("BAD-MONTH", bad.status_code == 400, bad.get_json().get("error", ""))
bad = c.put("/api/products/%d/report" % pid, json={"saas_months": {"2026-01": {"new": -1}}})
check("NEGATIVE", bad.status_code == 400, bad.get_json().get("error", ""))
report(projects={"999999": {"kind": "saas"}, str(lic["id"]): {"kind": "project"}})
conn = database.get_db(database.DEFAULT_DB)
import json  # noqa: E402
stored = json.loads(conn.execute("SELECT report FROM products WHERE id=?", (pid,)).fetchone()[0])
conn.close()
check("UNLINKED", "999999" not in stored["projects"], "settings for a project that is not linked are dropped")
from werkzeug.security import generate_password_hash  # noqa: E402
conn = database.get_db(database.DEFAULT_DB)
conn.execute("INSERT INTO users (username, password_hash, full_name, role, company_access) VALUES (?,?,?,?,?)",
             ("rview", generate_password_hash("viewer-pass-1"), "Report viewer", "viewer", "all"))
conn.commit()
conn.close()
v = server.app.test_client()
assert v.post("/api/login", json={"username": "rview", "password": "viewer-pass-1"}).status_code == 200
check("VIEWER", v.get("/api/products/%d/report?year=%d" % (pid, Y)).status_code == 200
      and v.put("/api/products/%d/report" % pid, json={"title": "x"}).status_code == 403,
      "a viewer reads the report but cannot change its settings")

print("=== the settings survive an Excel model import ===")
report(title="SENTIMIND", headcount=12, projects={str(nx["id"]): {"kind": "saas", "note": "keep me"}})
wb = c.get("/api/products/%d/workbook" % pid).data
r = c.post("/api/products/%d/import" % pid, data={"file": (io.BytesIO(wb), "sbu.xlsx")},
           content_type="multipart/form-data").get_json()
after_imp = c.get("/api/products/%d/report?year=%d" % (pid, Y)).get_json()["settings"]
check("SURVIVE", r.get("ok") and after_imp["title"] == "SENTIMIND" and after_imp["headcount"] == 12
      and after_imp["projects"][str(nx["id"])]["note"] == "keep me",
      "replacing the model from Excel leaves the report settings alone")

print("=== the workbook ===")
x = c.get("/api/products/%d/report.xlsx?year=%d" % (pid, Y))
book = load_workbook(io.BytesIO(x.data))
check("SHEETS", book.sheetnames == ["DASHBOARD", "P&L YTD", "CASH FLOW", "TRACKER PROYEK",
                                     "INDIKATOR KESEHATAN", "KINERJA SaaS"], ", ".join(book.sheetnames))
check("BANNER", str(book["DASHBOARD"]["B1"].value).startswith("SENTIMIND  |  FINANCIAL DASHBOARD EKSEKUTIF"),
      book["DASHBOARD"]["B1"].value)
pl = book["P&L YTD"]
labels = {str(pl.cell(row=r, column=2).value or "").strip(): r for r in range(1, pl.max_row + 1)}
tot = labels.get("TOTAL PENDAPATAN")
check("FORMULAS", tot and str(pl.cell(row=tot, column=40).value).startswith("="),
      "totals and YTD are live formulas: %s" % (pl.cell(row=tot, column=40).value if tot else "?"))
rep7 = c.get("/api/products/%d/report?year=%d" % (pid, Y)).get_json()
cf = book["CASH FLOW"]
n_rows = sum(1 for r in range(8, cf.max_row + 1) if cf.cell(row=r, column=2).value
             and str(cf.cell(row=r, column=2).value).startswith("JV"))
check("CASHROWS", n_rows == sum(1 for r in rep7["cash"]["rows"] if r["closed"]),
      "%d cash rows, the closed months only" % n_rows)

print()
print("RESULT: %d passed, %d failed" % (ok, fail))
shutil.rmtree(TMP, ignore_errors=True)
sys.exit(1 if fail else 0)
