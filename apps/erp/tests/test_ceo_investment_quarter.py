"""v1.11 batch 2: the CEO Dashboard, journal lines charged to an investment,
the main company, and the Oracle's quarterly NAME | KODE | VALUE sheet.

Covers the rules that must not quietly change: a line belongs to a project OR
an investment; an investment's "paid" includes every posted line charged to it;
the cash forecast starts from the last ACTUAL closing; and a quarterly sheet
refuses a positive value on a cost account rather than importing it.

    python apps/erp/tests/test_ceo_investment_quarter.py
"""
import datetime
import io
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
TMP = tempfile.mkdtemp(prefix="ceo_")
os.environ["MORES_HV_DATA_DIR"] = TMP

import database  # noqa: E402
database.create_database(database.DEFAULT_DB, seed_demo=True)
import fpa_cash  # noqa: E402
import server  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from werkzeug.security import generate_password_hash  # noqa: E402

ok = fail = 0


def check(name, cond, detail=""):
    global ok, fail
    if cond:
        ok += 1
        print("  PASS  %-12s %s" % (name, detail))
    else:
        fail += 1
        print("  FAIL  %-12s %s" % (name, detail))


c = server.app.test_client()
assert c.post("/api/login", json={"username": "admin", "password": "MoresMores2018"}).status_code == 200
companies = {x["code"]: x for x in c.get("/api/companies").get_json()}
mda = companies["MDA"]["id"]

print("=== the main company ===")
me = c.get("/api/me").get_json()
check("MAIN-MDA", me["main_company_id"] == mda, "MDA unless Settings says otherwise")
other = next(x["id"] for k, x in companies.items() if k != "MDA")
c.put("/api/settings/main-company", json={"company_id": other})
check("MAIN-SET", c.get("/api/me").get_json()["main_company_id"] == other, "an admin can name another")
c.put("/api/settings/main-company", json={"company_id": mda})
conn = database.get_db(database.DEFAULT_DB)
conn.execute("INSERT INTO users (username, password_hash, full_name, role, company_access) VALUES (?,?,?,?,?)",
             ("ceoview", generate_password_hash("viewer-pass-1"), "Viewer", "viewer", "all"))
conn.commit()
conn.close()
v = server.app.test_client()
assert v.post("/api/login", json={"username": "ceoview", "password": "viewer-pass-1"}).status_code == 200
check("MAIN-ADMIN", v.put("/api/settings/main-company", json={"company_id": other}).status_code == 403,
      "only an admin changes it")

print("=== a journal line charged to an investment ===")
accts = {a["code"]: a for a in c.get("/api/accounts?company_id=%d" % mda).get_json()}
direct, bank = accts["5100-01"]["id"], accts["1120"]["id"]
inv = c.post("/api/investments", json={"company_id": mda, "name": "Fellowship Program 2026",
                                       "committed_amount": 300000000}).get_json()["id"]
project = next(p for p in c.get("/api/projects?company_id=all").get_json() if p["company_id"] == mda)
line = lambda **kw: dict({"account_id": direct, "debit": 25000000, "credit": 0}, **kw)
post = lambda lines, date="2026-03-10": c.post("/api/journals", json={
    "company_id": mda, "date": date, "description": "fellowship honor", "status": "posted", "lines": lines})
r = post([line(project_id=project["id"], investment_id=inv), {"account_id": bank, "debit": 0, "credit": 25000000}])
check("NOT-BOTH", r.status_code == 400, r.get_json().get("error", ""))
r = post([line(investment_id=999999), {"account_id": bank, "debit": 0, "credit": 25000000}])
check("UNKNOWN", r.status_code == 400, "an investment that does not exist is refused")
r = post([line(investment_id=inv), {"account_id": bank, "debit": 0, "credit": 25000000}])
check("POSTED", r.status_code == 201, "Direct cost 5100-01 charged to the investment")
jid = r.get_json()["id"]
jl = c.get("/api/journals/%d" % jid).get_json()["lines"]
check("SHOWN", any(l.get("investment_name") == "Fellowship Program 2026" and not l.get("project_code") for l in jl),
      "the entry shows the investment, not a project")
c.post("/api/investments/%d/events" % inv, json={"kind": "outflow", "date": "2026-02-01",
                                                  "description": "by hand", "amount": 10000000})
row = next(x for x in c.get("/api/investments?company_id=all").get_json() if x["id"] == inv)
check("PAID", row["invested"] == 35000000 and row["invested_ledger"] == 25000000 and row["invested_entered"] == 10000000,
      "paid = 10 jt entered + 25 jt from the ledger")
det = c.get("/api/investments/%d" % inv).get_json()
check("DETAIL", len(det["ledger_lines"]) == 1 and det["ledger_lines"][0]["account_code"] == "5100-01",
      "the investment lists the line charged to it")
draft = c.post("/api/journals", json={"company_id": mda, "date": "2026-03-11", "description": "draft", "status": "draft",
                                      "lines": [line(investment_id=inv), {"account_id": bank, "debit": 0, "credit": 25000000}]})
row = next(x for x in c.get("/api/investments?company_id=all").get_json() if x["id"] == inv)
check("POSTED-ONLY", row["invested_ledger"] == 25000000, "a draft does not count")
r = c.delete("/api/investments/%d" % inv)
check("GUARD", r.status_code == 400, "an investment with lines charged to it cannot be deleted")

print("=== the CEO dashboard ===")
c.post("/api/investments/%d/hot" % inv, json={"is_hot": True})
d = c.get("/api/dashboard/ceo?company_id=all&year=2026").get_json()
check("PARTS", all("error" not in (d[k] if isinstance(d[k], dict) else {}) for k in ("cash", "highlights", "investments"))
      and isinstance(d["sbu"], list), "cash, highlights, SBU and investments all answer")
cash = d["cash"]
T = cash["through"]
conn = database.get_db(database.DEFAULT_DB)
ids = [x["id"] for x in companies.values()]
want = fpa_cash.opening_cash(conn, ids, "2026-%02d-01" % (T + 1) if T < 12 else "2027-01-01", None)
conn.close()
check("ACTUAL", T >= 1 and abs(cash["months"][T - 1]["actual"] - want) < 1,
      "closing cash of %s = the ledger's cash on the 1st after (%s)" % (cash["through_label"], want))
check("ANCHOR", cash["months"][T - 1]["forecast"] == cash["months"][T - 1]["actual"]
      and all(cash["months"][m]["actual"] is None for m in range(T, 12)), "the forecast starts at the last actual closing")
steps = all(abs(cash["months"][m]["forecast"] - cash["months"][m - 1]["forecast"] - cash["months"][m]["budget_net"]) < 1
            for m in range(T, 12))
check("FORECAST", steps, "each forecast month adds that month's budgeted net")
check("MINIMUM", abs(cash["minimum"] - sum(f["floor"] for f in cash["floors"])) < 1
      and cash["status"] == ("above" if cash["latest"] >= cash["minimum"] else "below"),
      "minimum = the Oracle's floors, status %s" % cash["status"])
check("HOT-TOP", d["investments"]["rows"][0]["id"] == inv and d["investments"]["rows"][0]["is_hot"],
      "the starred investment is first")

# highlights: a track stuck in SPM for 40 days against a 14-day plan
tid = c.post("/api/money-tracker", json={"company_id": mda, "title": "Termin 2 Dinas", "amount": 400000000,
                                         "phase_key": "p7", "status": "active",
                                         "started_at": (datetime.date.today() - datetime.timedelta(days=40)).isoformat()}
             ).get_json()["id"]
done = c.post("/api/money-tracker", json={"company_id": mda, "title": "Termin 1 Dinas", "amount": 300000000,
                                          "phase_key": "p12", "status": "done"}).get_json()["id"]
h = c.get("/api/dashboard/ceo?company_id=all&year=2026").get_json()["highlights"]
late = next((r for r in h["delayed"]["rows"] if r["id"] == tid), None)
check("DELAYED", late is not None and late["days_late"] == 26, "40 days in a 14-day phase = 26 days late")
check("SPM", any(r["id"] == tid for r in h["spm"]["rows"]), "and it is in the SPM-to-SP2D stretch")
check("DONE", any(r["id"] == done for r in h["done"]["rows"]), "a paid track is done")

print("=== the Oracle's quarterly sheet ===")
x = c.get("/api/templates/oracle-quarter?company_id=%d&year=2026&quarter=4" % mda)
wb = load_workbook(io.BytesIO(x.data))
ws = wb["Quarter plan"]
heads = [ws.cell(row=3, column=col).value for col in range(1, ws.max_column + 1) if ws.cell(row=3, column=col).value]
check("LAYOUT", x.status_code == 200 and heads.count("W1") == 3 and ws.cell(row=4, column=3).value == "OB"
      and ws.cell(row=4, column=4).value == "KODE", "3 months x W1..W4, OB, and a KODE column per week")
conn = database.get_db(database.DEFAULT_DB)
budget_rows = conn.execute(
    "SELECT COUNT(*) FROM (SELECT 1 FROM budgets b JOIN accounts a ON a.id=b.account_id WHERE b.company_id=?"
    " AND b.year=2026 AND b.month IN (10,11,12) AND a.type IN ('revenue','expense')"
    " GROUP BY b.month, b.week, a.code, a.type, b.project_id HAVING ABS(SUM(b.amount)) > 0.004)", (mda,)).fetchone()[0]
budget_in = conn.execute(
    "SELECT COALESCE(SUM(b.amount),0) FROM budgets b JOIN accounts a ON a.id=b.account_id WHERE b.company_id=?"
    " AND b.year=2026 AND b.month IN (10,11,12) AND a.type='revenue'", (mda,)).fetchone()[0]
conn.close()


def upload(data, preview=True, name=None):
    form = {"file": (io.BytesIO(data), "q.xlsx"), "company_id": str(mda), "year": "2026"}
    if preview:
        form["preview"] = "1"
    if name:
        form["name"] = name
    return c.post("/api/plan/import-cashflow", data=form, content_type="multipart/form-data").get_json()


back = upload(x.data)
check("ROUNDTRIP", back["ok"] and back["count"] == budget_rows and back["coded"] == budget_rows
      and abs(back["total_in"] - budget_in) < 1,
      "the downloaded sheet imports back: %d rows, every one on its KODE" % back["count"])


def sheet(rows):
    """A hand-kept sheet with the KODE column: (label, code, value) in W1 Oct."""
    wbk = Workbook()
    s = wbk.active
    s["C2"] = datetime.datetime(2026, 10, 1)
    for i, w in enumerate(("W1", "W2", "W3", "W4")):
        s.cell(row=3, column=3 + 3 * i, value=w)
        s.merge_cells(start_row=3, start_column=3 + 3 * i, end_row=3, end_column=5 + 3 * i)
        s.cell(row=4, column=4 + 3 * i, value="KODE")
    s["C4"], s["E4"] = "OB", 500000000
    for n, (label, code, value) in enumerate(rows):
        s.cell(row=5 + n, column=3, value=label)
        s.cell(row=5 + n, column=4, value=code)
        s.cell(row=5 + n, column=5, value=value)
    buf = io.BytesIO()
    wbk.save(buf)
    return buf.getvalue()


rev_code = next(a["code"] for a in accts.values() if a["type"] == "revenue")
bad = upload(sheet([("Medmon Okt", "5100-01", 100000000)]))
check("SIGN", not bad["ok"] and "cost account" in bad["errors"][0], bad["errors"][0][:90] if bad.get("errors") else "")
bad = upload(sheet([("Medmon Okt", "9999-99", 100000000)]))
check("NO-ACCOUNT", not bad["ok"] and "9999-99" in bad["errors"][0], "a KODE that is not an account is refused")
good = upload(sheet([(project["name"], rev_code, 150000000), ("Fellowship Program 2026", "5100-01", -30000000),
                     ("Server fisik", "", -20000000)]), preview=False, name="Q4 coded")
check("CODED", good["ok"] and good["coded"] == 2 and good["tagged"] == 2,
      "2 rows on their KODE, 2 tagged (a project and an investment), 1 on the default account")
conn = database.get_db(database.DEFAULT_DB)
pw = {r["note"]: dict(r) for r in conn.execute(
    "SELECT w.note, w.project_id, w.investment_id, a.code FROM plan_weeks w JOIN accounts a ON a.id=w.account_id"
    " WHERE w.version_id=?", (good["version_id"],))}
conn.close()
check("STORED", pw[project["name"]]["project_id"] == project["id"] and pw[project["name"]]["code"] == rev_code
      and pw["Fellowship Program 2026"]["investment_id"] == inv and pw["Fellowship Program 2026"]["code"] == "5100-01"
      and pw["Server fisik"]["code"] in ("6900", "6600"),
      "the scenario keeps the account, the project and the investment")

print()
print("RESULT: %d passed, %d failed" % (ok, fail))
shutil.rmtree(TMP, ignore_errors=True)
sys.exit(1 if fail else 0)
