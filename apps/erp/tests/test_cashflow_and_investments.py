"""v1.09: the cashflow workbook the Oracle reads, and investment progress.

Covers the rules that must not quietly change: a hand-kept cash sheet is read
the way finance writes it, a blank row ends the plan, a preview writes nothing,
and progress is weighted by milestone rather than by money spent.

    python apps/erp/tests/test_cashflow_and_investments.py
"""
import datetime
import io
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
TMP = tempfile.mkdtemp(prefix="cf_inv_")
os.environ["MORES_HV_DATA_DIR"] = TMP

import database  # noqa: E402
database.create_database(database.DEFAULT_DB, seed_demo=True)
import excel_io as X  # noqa: E402
import server  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

ok = fail = 0


def check(name, cond, detail=""):
    global ok, fail
    if cond:
        ok += 1
        print("  PASS  %-12s %s" % (name, detail))
    else:
        fail += 1
        print("  FAIL  %-12s %s" % (name, detail))


def sheet(opening=823366000.0, break_it=False, months=2):
    """A cash sheet shaped like the one finance keeps: months across the top,
    W1..W4 under each, label + amount pairs below, then a scratch block that a
    blank row separates from the plan."""
    wb = Workbook()
    ws = wb.active
    cols = {}
    c = 3
    for mi in range(months):
        ws.cell(row=2, column=c, value=datetime.datetime(2026, 9 + mi, 1))
        for w in range(1, 5):
            ws.cell(row=3, column=c, value="W%d" % w)
            cols[(9 + mi, w)] = c
            c += 2
        c += 1                                   # a spacer column between months
    ws.cell(row=4, column=cols[(9, 4)], value="OB")
    ws.cell(row=4, column=cols[(9, 4)] + 1, value=opening)
    rows = {
        (9, 4): [("JKT SLR", -183000000), ("Office Allocation JKT", -20000000)],
        (10, 1): [("CDA Sept", 97000000), ("Survei WKS Term 1", -100000000), ("No amount here", None)],
        (10, 2): [("Outlook Pariwisata", 172000000)],
    }
    for key, items in rows.items():
        for i, (label, amt) in enumerate(items):
            ws.cell(row=5 + i, column=cols[key], value=label)
            if amt is not None:
                ws.cell(row=5 + i, column=cols[key] + 1, value=amt)
    # the sheet's own week openings, which must agree with the rows above
    run = opening
    for key in [(9, 4), (10, 1), (10, 2), (10, 3)]:
        if key != (9, 4):
            ws.cell(row=4, column=cols[key] + 1, value=run + (7 if break_it and key == (10, 1) else 0))
        run += sum(a for _, a in rows.get(key, []) if a)
    ws.cell(row=12, column=cols[(9, 4)] + 1, value=run)      # closing row, no labels
    # blank row 13, then somebody's scratch working that must NOT be imported
    ws.cell(row=14, column=cols[(9, 1)], value="C BDKR")
    ws.cell(row=14, column=cols[(9, 1)] + 1, value=2481595787)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


print("=== reading a hand-kept cash sheet ===")
plan, errors = X.parse_cashflow_workbook(io.BytesIO(sheet()), 2026)
check("PARSE", plan is not None and not errors, str(errors))
by = {b["label"]: b for b in plan["weeks"]}
check("WEEKS", sorted(by) == ["W1 Oct 2026", "W2 Oct 2026", "W3 Oct 2026", "W4 Sep 2026"], str(sorted(by)))
check("SIGN", [(i["label"], i["flow"], i["amount"]) for i in by["W1 Oct 2026"]["items"]]
      == [("CDA Sept", "in", 97000000.0), ("Survei WKS Term 1", "out", 100000000.0)],
      "positive is money in, negative is money out")
check("OPENING", by["W4 Sep 2026"]["opening"] == 823366000.0, "the OB row is the opening balance")
check("SCRATCH", not any("C BDKR" in i["label"] for b in plan["weeks"] for i in b["items"])
      and any("C BDKR" in x for x in plan["ignored"]),
      "the block under the blank row is left out, and said so")
check("NO-AMOUNT", any("No amount here" in w for w in plan["warnings"]),
      "a label with no amount beside it is reported, not guessed at")
check("RECONCILE", plan["reconciled"] is True and not plan["checks"],
      "each week's opening plus its rows equals the next week's opening")
check("TOTALS", (plan["total_in"], plan["total_out"]) == (269000000.0, 303000000.0),
      "in %s / out %s" % (plan["total_in"], plan["total_out"]))

bad, _ = X.parse_cashflow_workbook(io.BytesIO(sheet(break_it=True)), 2026)
check("MISMATCH", bad["reconciled"] is False and bad["checks"],
      "one edited opening balance is caught (it puts %d weeks out): %s"
      % (len(bad["checks"]), bad["checks"][0][:64] if bad["checks"] else ""))

empty = Workbook()
empty.active["A1"] = "no week headers here"
buf = io.BytesIO()
empty.save(buf)
none, errs = X.parse_cashflow_workbook(io.BytesIO(buf.getvalue()), 2026)
check("NOT-A-PLAN", none is None and "W1" in errs[0], errs[0] if errs else "")
none, errs = X.parse_cashflow_workbook(io.BytesIO(b"not a workbook"), 2026)
check("NOT-XLSX", none is None and "could not be opened" in errs[0], "")

c = server.app.test_client()
assert c.post("/api/login", json={"username": "admin", "password": "MoresMores2018"}).status_code == 200
companies = c.get("/api/companies").get_json()
cid = companies[0]["id"]


def upload(url, data, fields):
    form = {"file": (io.BytesIO(data), "cash.xlsx")}
    form.update(fields)
    return c.post(url, data=form, content_type="multipart/form-data")


def versions():
    return c.get("/api/plan/versions").get_json()["versions"]


print("=== importing it as an Oracle scenario ===")
before = len(versions())
r = upload("/api/plan/import-cashflow", sheet(), {"company_id": str(cid), "year": "2026", "preview": "1"})
prev = r.get_json()
check("PREVIEW", prev.get("ok") and prev["count"] == 5 and len(versions()) == before,
      "a preview reports %d items and writes nothing" % prev.get("count", -1))
check("PREV-ACCT", "in" in prev.get("accounts", {}) and "out" in prev["accounts"],
      "it says which accounts the money will book to: %s" % prev.get("accounts"))
check("PREV-OPEN", any("ledger" in w for w in prev["warnings"]),
      "it says the Oracle starts from the ledger's cash, not the sheet's opening")

r = upload("/api/plan/import-cashflow", sheet(),
           {"company_id": str(cid), "year": "2026", "name": "Q4 drive",
            "certainty_in": "expected", "certainty_out": "committed"})
res = r.get_json()
check("IMPORT", r.status_code == 201 and res.get("ok") and len(versions()) == before + 1,
      "scenario '%s' created" % res.get("name"))
vid = res.get("version_id")
conn = database.get_db(database.DEFAULT_DB)
rows = [dict(x) for x in conn.execute("SELECT * FROM plan_weeks WHERE version_id=?", (vid,))]
check("ROWS", len(rows) == 5, "%d plan rows written" % len(rows))
check("ROW-FLOW", sorted((r["month"], r["week"], r["flow"], r["amount"]) for r in rows)
      == [(9, 4, "out", 20000000.0), (9, 4, "out", 183000000.0), (10, 1, "in", 97000000.0),
          (10, 1, "out", 100000000.0), (10, 2, "in", 172000000.0)],
      "month, week, direction and amount all survive the trip")
check("ROW-CERT", {r["certainty"] for r in rows if r["flow"] == "in"} == {"expected"}
      and {r["certainty"] for r in rows if r["flow"] == "out"} == {"committed"},
      "each direction keeps the certainty it was imported with")
check("ROW-NOTE", any(r["note"] == "Outlook Pariwisata" for r in rows),
      "the sheet's own wording is kept on the row")
conn.close()

d = c.post("/api/oracle/consult", json={"version_id": vid, "year": 2026}).get_json()
live = c.post("/api/oracle/consult", json={"version_id": None, "year": 2026}).get_json()
# the count also carries receivables, payables and investment commitments, so
# what matters is that the scenario's own rows are in there and the live budget's are not
check("ORACLE", d.get("items", 0) >= 5 and d.get("items") != live.get("items")
      and d.get("verdict") in ("KRITIS", "TOLAK", "WASPADA", "LULUS"),
      "the Oracle reads the imported scenario: %s on %d dated items (live budget: %d)"
      % (d.get("verdict"), d.get("items", 0), live.get("items", 0)))
check("ORACLE-WK", len(((d.get("entities") or [{}])[0]).get("weekly") or []) == 48,
      "48 weekly buckets for the chart")

print("=== the Oracle's own plan template ===")
r = c.get("/api/templates/oracle-plan?company_id=%d&year=2026" % cid)
check("TPL-OK", r.status_code == 200 and r.data[:2] == b"PK", "the template downloads")
tpl = r.data
wb = load_workbook(io.BytesIO(tpl))
check("TPL-SHEETS", wb.sheetnames == ["Cash plan", "Accounts", "Projects", "Project revenue & budget"],
      str(wb.sheetnames))
ws = wb["Cash plan"]
body = [[ws.cell(row=i, column=j).value for j in range(1, 12)]
        for i in range(5, ws.max_row + 1) if ws.cell(row=i, column=1).value]
check("TPL-BUDGET", len(body) > 0 and all(b[3] in X.MONTHS and str(b[4]).startswith("W") for b in body),
      "the budget already in the database is written out to scroll: %d rows" % len(body))
dvs = " ".join(str(d.formula1) for d in ws.data_validations.dataValidation)
check("TPL-DROPDOWN", "Accounts!" in dvs and "Projects!" in dvs and "committed" in dvs and "W1,W2" in dvs,
      "account, project, week and certainty are drop-downs, not free text")
accounts = {wb["Accounts"].cell(row=i, column=1).value: wb["Accounts"].cell(row=i, column=3).value
            for i in range(4, wb["Accounts"].max_row + 1) if wb["Accounts"].cell(row=i, column=1).value}
COST = next(code for code, kind in accounts.items() if kind == "expense" and code.startswith("5100"))
REV = next(code for code, kind in accounts.items() if kind == "revenue")
prj = wb["Projects"].cell(row=4, column=1).value
check("TPL-LISTS", len(accounts) > 5 and prj, "%d accounts and project %s to pick from" % (len(accounts), prj))
stats = wb["Project revenue & budget"]
srows = [[stats.cell(row=i, column=j).value for j in range(1, 8)]
         for i in range(5, stats.max_row + 1) if stats.cell(row=i, column=1).value]
check("TPL-PRJ-REV", srows and any(x[2] or x[3] for x in srows),
      "each project's revenue and budget is there for context: %s" % (srows[0][:4] if srows else ""))


def plan_book(rows, sheet="Cash plan"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["Oracle cash plan"])
    ws.append([])
    ws.append([])
    ws.append(X.PLAN_HEADERS)
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# the line the user asked to be able to write, straight from the drop-downs
EXAMPLE = [COST, prj, 2026, "Dec", "W2", "out", 25000000, "planned", "operating", "Termin design"]
r = upload("/api/plan/import-cashflow", plan_book([EXAMPLE]),
           {"company_id": str(cid), "year": "2026", "preview": "1"})
pv = r.get_json()
check("PLAN-PREVIEW", pv.get("ok") and pv["count"] == 1 and pv["kind"] == "plan"
      and pv["weeks"][0]["label"] == "W2 Dec 2026" and pv["total_out"] == 25000000,
      "'%s on %s, 25 jt, out, W2 December 2026' reads back as %s"
      % (COST, prj, pv.get("weeks", [{}])[0].get("label")))
check("PLAN-BYACC", pv["by_account"][0]["code"] == COST and pv["by_account"][0]["project"] == prj,
      "and the preview says which account and project it lands on")

before = len(versions())
r = upload("/api/plan/import-cashflow", plan_book([EXAMPLE, [REV, "", 2026, "Mar", "W1", "", 900000000, "", "", "Termin 1"]]),
           {"company_id": str(cid), "year": "2026", "name": "Plan from template"})
res = r.get_json()
check("PLAN-IMPORT", r.status_code == 201 and res.get("ok") and len(versions()) == before + 1,
      "scenario '%s' with %d rows" % (res.get("name"), res.get("count", 0)))
conn = database.get_db(database.DEFAULT_DB)
rows = [dict(x) for x in conn.execute(
    "SELECT w.*, a.code AS acode, p.code AS pcode FROM plan_weeks w"
    " JOIN accounts a ON a.id=w.account_id LEFT JOIN projects p ON p.id=w.project_id"
    " WHERE w.version_id=? ORDER BY w.month", (res["version_id"],))]
conn.close()
ex = next((x for x in rows if x["acode"] == COST), None)
check("PLAN-ROW", ex and (ex["month"], ex["week"], ex["flow"], ex["amount"], ex["pcode"],
                          ex["certainty"], ex["cf_class"]) == (12, 2, "out", 25000000.0, prj,
                                                               "planned", "operating"),
      "the row keeps its account, project, week, direction and certainty")
rev = next((x for x in rows if x["acode"] == REV), None)
check("PLAN-DERIVED", rev and rev["flow"] == "in" and rev["project_id"] is None
      and rev["certainty"] == "planned",
      "a blank Direction follows the account type (revenue = money in), blank certainty takes the default")

bad = upload("/api/plan/import-cashflow", plan_book([
    ["9999", "", 2026, "Dec", "W2", "out", 1000, "", "", ""],
    [COST, "NO-SUCH", 2026, "Dec", "W2", "out", 1000, "", "", ""],
    [COST, "", 2026, "Smarch", "W2", "out", 1000, "", "", ""],
    [COST, "", 2026, "Dec", "W9", "out", 1000, "", "", ""],
    [COST, "", 2026, "Dec", "W2", "out", 1000, "maybe", "", ""]]),
    {"company_id": str(cid), "year": "2026"}).get_json()
check("PLAN-ERRORS", bad.get("ok") is False and len(bad["errors"]) == 5
      and all(e.startswith("Row ") for e in bad["errors"]),
      "every bad row is named with its row number: %s" % "; ".join(bad["errors"])[:110])
check("PLAN-ATOMIC", len(versions()) == before + 1, "and nothing was written")

full = load_workbook(io.BytesIO(tpl))
fs = full["Cash plan"]
nxt = 5 + len(body)
for j, v in enumerate(EXAMPLE, start=1):
    fs.cell(row=nxt, column=j, value=v)
buf = io.BytesIO()
full.save(buf)
res = upload("/api/plan/import-cashflow", buf.getvalue(),
             {"company_id": str(cid), "year": "2026", "name": "Budget plus one line"}).get_json()
check("PLAN-ROUNDTRIP", res.get("ok") and res["count"] == len(body) + 1,
      "the downloaded template imports back whole: %d budget rows + the new line" % len(body))

print("=== investment progress: milestones, not money spent ===")
iid = c.post("/api/investments", json={"company_id": cid, "name": "Test initiative",
                                       "committed_amount": 500000000}).get_json()["id"]
ids = [c.post("/api/investments/%d/milestones" % iid,
              json={"title": t, "weight": w, "due_date": due}).get_json()["id"]
       for t, w, due in [("Pilot signed off", 1, "2026-03-31"),
                         ("Rolled out", 3, "2026-09-30"),
                         ("Reviewed", 1, "")]]
inv = c.get("/api/investments/%d" % iid).get_json()
check("MS-ADD", len(inv["milestones"]) == 3 and inv["progress"]["pct"] == 0,
      "three milestones, nothing done yet")
c.put("/api/investments/%d/milestones/%d" % (iid, ids[1]), json={"done": True})
inv = c.get("/api/investments/%d" % iid).get_json()
check("MS-WEIGHT", inv["progress"]["pct"] == 60.0 and inv["progress"]["done"] == 1,
      "one of three done, but it carries 3 of 5 weight -> 60%, not 33%")
c.put("/api/investments/%d/milestones/%d" % (iid, ids[1]), json={"done": False})
check("MS-UNDO", c.get("/api/investments/%d" % iid).get_json()["progress"]["pct"] == 0.0,
      "un-ticking puts it back")
c.delete("/api/investments/%d/milestones/%d" % (iid, ids[2]))
inv = c.get("/api/investments/%d" % iid).get_json()
check("MS-DEL", len(inv["milestones"]) == 2, "a milestone can be removed")
check("MS-BAD", c.post("/api/investments/%d/milestones" % iid, json={"title": "  "}).status_code == 400,
      "a milestone needs a name")

print("=== an investment that IS a project ===")
pid = c.get("/api/projects?company_id=%d" % cid).get_json()[0]["id"]
c.put("/api/investments/%d" % iid, json={"name": "Test initiative", "company_id": cid,
                                         "committed_amount": 500000000, "linked_project_id": pid})
c.post("/api/investments/%d/commitments" % iid,
       json={"year": 2026, "month": 11, "week": 2, "amount": 200000000, "note": "Termin 1"})
link = c.get("/api/projects/%d/investment" % pid).get_json()["investment"]
check("PRJ-INV", link and link["id"] == iid and link["name"] == "Test initiative",
      "the project knows which investment funds it")
check("PRJ-CASH", [(x["month"], x["week"], x["amount"], x["note"]) for x in link["commitments"]]
      == [(11, 2, 200000000.0, "Termin 1")], "with the payment schedule the Oracle counts")
check("PRJ-UNSCHED", link["unscheduled"] == 300000000.0,
      "and what is committed but still unscheduled (500 jt - 200 jt)")
check("PRJ-PROG", link["progress"]["milestones"] == 2, "and its progress")
plist = c.get("/api/projects?company_id=%d" % cid).get_json()
check("PRJ-BADGE", any(p["id"] == pid and p["investment_id"] == iid for p in plist),
      "the projects list carries the link, for the (INVESTMENT) badge")
other = [p for p in plist if p["id"] != pid]
check("PRJ-PLAIN", all(p["investment_id"] is None for p in other) if other else True,
      "ordinary projects are not labelled")

print("=== contract value, and COGS kept apart from operating expense ===")
perf = c.get("/api/projects/performance?company_id=all&year=2026").get_json()
prow = next((x for x in perf["rows"] if x["expense"]), None)
check("PERF-SPLIT", prow and round(prow["cogs"] + prow["opex"], 2) == round(prow["expense"], 2),
      "COGS %s + opex %s = the project's whole expense" % (prow["cogs"], prow["opex"]) if prow else "")
check("PERF-BY", prow and abs(sum(a["amount"] for a in prow["cogs_by"].values()) - prow["cogs"]) < 0.01
      and all(str(a["code"]).startswith("5") for a in prow["cogs_by"].values()),
      "the breakdown is by 5xxx account and adds up: %s" % sorted(prow["cogs_by"]) if prow else "")
# the bug this guards: on a real chart of accounts the project costs sit on 6000
# accounts, so counting only the 5000 family reported no cost and a 100% margin
adds_up = all(abs(sum(a["amount"] for a in x["cost_by"].values()) - x["expense"]) < 0.01
              for x in perf["rows"])
check("PERF-COST", prow and adds_up and len(prow["cost_by"]) >= len(prow["cogs_by"]),
      "cost_by carries EVERY expense account on the project and adds up to its expense")
pid6 = prow["project_id"]
p6 = next(x for x in c.get("/api/projects?company_id=all").get_json() if x["id"] == pid6)
accs = c.get("/api/accounts?company_id=%d" % p6["company_id"]).get_json()
mkt = next(a for a in accs if a["code"].startswith("6") and a["type"] == "expense")
cash = next(a for a in accs if a["code"].startswith("11") and a["type"] == "asset")
before6 = prow["expense"]
r6 = c.post("/api/journals", json={
    "company_id": p6["company_id"], "date": "2026-04-15", "status": "posted",
    "description": "Marketing booked straight to the project (6000 family)",
    "lines": [{"account_id": mkt["id"], "project_id": pid6, "debit": 120000000, "credit": 0},
              {"account_id": cash["id"], "debit": 0, "credit": 120000000}]})
after = next(x for x in c.get("/api/projects/performance?company_id=all&year=2026").get_json()["rows"]
             if x["project_id"] == pid6)
check("PERF-COST-6", r6.status_code == 201
      and round(after["expense"] - before6, 2) == 120000000.0
      and mkt["code"] in after["cost_by"]
      and mkt["code"] not in after["cogs_by"],
      "a 6000-family cost tagged to the project counts as its cost (%s), which is what a real "
      "ledger books and what reported a 100%% margin before" % mkt["code"])
p = next(x for x in c.get("/api/projects?company_id=all").get_json() if x["id"] == prow["project_id"])
check("CONTRACT-0", p["contract_value"] == 0, "a project starts with no contract value")
c.put("/api/projects/%d" % p["id"], json={"name": p["name"], "status": p["status"],
                                          "contract_value": 1750000000})
again = next(x for x in c.get("/api/projects?company_id=all").get_json() if x["id"] == p["id"])
check("CONTRACT", again["contract_value"] == 1750000000 and again["name"] == p["name"],
      "the contract value is stored on the project, and saving it changes nothing else")

print("=== menus keep up with new sections ===")
conn = database.get_db(database.DEFAULT_DB)
conn.execute("INSERT INTO users (username, password_hash, full_name, role, company_access,"
             " menu_access) VALUES ('bod2','x','Board','viewer','all','dashboard,reports')")
conn.commit()
database.migrate_database(conn)
got = conn.execute("SELECT menu_access FROM users WHERE username='bod2'").fetchone()[0]
check("MENU-FIX", set(got.split(",")) == {"dashboard", "reports", "oracle", "product", "expenses"}, got)
allm = conn.execute("SELECT menu_access FROM users WHERE username='admin'").fetchone()[0]
check("MENU-ALL", allm == "all", "'all' is left alone")
conn.close()

print()
print("RESULT: %d passed, %d failed" % (ok, fail))
shutil.rmtree(TMP, ignore_errors=True)
sys.exit(1 if fail else 0)
